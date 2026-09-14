import base64
import copy
import io
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from grantthread.api import Binary, dispatch
from grantthread.domain import calculate
from grantthread.errors import DomainError
from grantthread.finance_math import convert_minor, normalize_rate
from grantthread.finance_service import FinancialService
from grantthread.repository import SQLiteRepository
from grantthread.seed import IDENTITIES, seed_all
from grantthread.service import Service
from grantthread.storage import LocalStorage


class Financials(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.env = patch.dict(os.environ, {'GRANTTHREAD_MODE': 'local', 'GRANTTHREAD_DATA_DIR': self.temp.name, 'BEDROCK_MODEL_ID': ''})
        self.env.start()
        self.repo = SQLiteRepository(Path(self.temp.name) / 'test.sqlite3')
        self.storage = LocalStorage(self.temp.name)
        seed_all(self.repo, self.storage)
        self.service = FinancialService(IDENTITIES['brightpath'], self.repo, self.storage)
        self.grant = self.service.create_grant({'name': 'Fictional international workshop', 'funderName': 'Fictional Cedar Trust', 'currency': 'CAD', 'award': '10000'})

    def tearDown(self):
        self.env.stop()
        self.temp.cleanup()

    def receipt(self, amount='1000', rate='0.3', received='2026-04-01', **extra):
        return self.service.add_receipt({'grantId': self.grant['id'], 'sourceCurrency': 'RON', 'receivedDate': received,
                                        'sourceAmount': amount, 'rate': rate, **extra})

    def draft(self, **extra):
        return self.service.create_entry({'grantId': self.grant['id'], 'kind': 'expense', 'date': '2026-04-15',
            'description': 'Fictional workshop supplies', 'reference': 'DEMO-INV-1', 'currency': 'RON', 'amount': '100',
            'category': '1.4 Goods and Supplies', 'conversionMode': 'manual', 'manualRate': '0.3', **extra})

    def confirm(self, entry):
        return self.service.confirm_entries({'entries': [{'id': entry['id'], 'expectedVersion': entry['version']}]})['entries'][0]

    def deny(self, action, code=None):
        with self.assertRaises(DomainError) as caught:
            action()
        if code:
            self.assertEqual(caught.exception.code, code)

    def workbook(self):
        from openpyxl import Workbook
        from datetime import datetime
        wb = Workbook()
        ws = wb.active
        ws.append(['Date', 'Description', 'Reference', 'Amount', 'Currency', 'Category', 'Rate'])
        ws.append([datetime(2026, 4, 15), 'Fictional workbook supplies', 'DEMO-XLSX-1', 100, 'RON', '1.4 Goods and Supplies', .321549266958621])
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    def test_receipt_direction_and_frozen_conversion(self):
        receipt = self.receipt(rate='4', rateDirection='source_per_report')
        self.assertEqual(receipt['rate'], '0.25')
        self.assertEqual(receipt['reportAmountMinor'], 25000)
        entry = self.confirm(self.draft(conversionMode='receipt', receiptId=receipt['id']))
        self.assertEqual(entry['reportAmountMinor'], 2500)
        self.receipt(amount='2000', rate='0.5')
        stored = self.service.overview()['entries'][0]
        self.assertEqual(stored['fxRate'], '0.25')
        self.assertEqual(stored['reportAmountMinor'], 2500)

    def test_financial_source_descriptor_preserves_grantee_scope_without_reading_bytes(self):
        result = self.service.upload_import({'kind': 'ledger', 'grantId': self.grant['id'], 'name': 'fictional-ledger.xlsx',
            'contentBase64': base64.b64encode(self.workbook()).decode()})
        path = f'/financials/imports/{result["id"]}/source'
        original = dispatch('GET', path, {}, IDENTITIES['brightpath'], self.repo, self.storage)
        self.assertIsInstance(original, Binary)
        self.assertTrue(original.data.startswith(b'PK'))
        with patch.object(self.storage, 'presign_get', return_value='https://synthetic-bucket.example.test/private-ledger') as sign, \
             patch.object(self.storage, 'get') as read_bytes:
            descriptor = dispatch('GET', path, {}, IDENTITIES['brightpath'], self.repo, self.storage)
            self.assertEqual(descriptor, {'downloadUrl': 'https://synthetic-bucket.example.test/private-ledger',
                                         'contentType': original.content_type, 'filename': original.name})
            self.assertTrue(sign.call_args.args[0].startswith('brightpath/'))
            sign.reset_mock()
            self.deny(lambda: dispatch('GET', path, {}, IDENTITIES['harbour'], self.repo, self.storage), 'not_found')
            self.deny(lambda: dispatch('GET', path, {}, IDENTITIES['northstar'], self.repo, self.storage), 'forbidden')
            sign.assert_not_called()
            read_bytes.assert_not_called()

    def test_weighted_rate_not_average_of_rates_and_excludes_future_receipts(self):
        first = self.receipt(amount='1000', rate='0.3')
        second = self.receipt(amount='3000', rate='0.4', received='2026-04-02')
        self.receipt(amount='9000', rate='0.9', received='2026-05-01')
        entry = self.confirm(self.draft(conversionMode='weighted_average'))
        self.assertEqual(entry['fxRate'], '0.375')
        self.assertEqual(entry['reportAmountMinor'], 3750)
        self.assertEqual(entry['conversion']['receiptIds'], sorted([first['id'], second['id']]))

    def test_automatic_receipt_uses_latest_eligible_marked_rate_and_freezes(self):
        pending = self.draft(conversionMode='receipt')
        self.assertIsNone(pending['reportAmountMinor'])
        first = self.receipt(useAsDefault=True)
        self.receipt(rate='0.9', received='2026-05-01', useAsDefault=True)
        self.receipt(rate='0.8', received='2026-04-03', useAsDefault=False)
        self.deny(lambda: self.confirm(pending), 'stale')
        pending = next(e for e in self.service.overview()['entries'] if e['id'] == pending['id'])
        confirmed = self.confirm(pending)
        self.assertEqual(confirmed['reportAmountMinor'], 3000)
        self.assertEqual(confirmed['conversion']['receiptId'], first['id'])
        self.assertTrue(confirmed['conversion']['automaticReceipt'])
        self.receipt(rate='0.4', received='2026-04-10', useAsDefault=True)
        self.assertEqual(next(e for e in self.service.overview()['entries'] if e['id'] == confirmed['id'])['reportAmountMinor'], 3000)
        newer = self.confirm(self.draft(conversionMode='receipt', reference='AUTO-NEW'))
        self.assertEqual(newer['reportAmountMinor'], 4000)
        from openpyxl import load_workbook
        exported = load_workbook(io.BytesIO(self.service.export_report(self.grant['id'])))
        self.assertIn(first['id'], exported['Ledger']['L4'].value)
        self.assertEqual(exported['Funding Receipts']['N4'].value, first['id'])

    def test_missing_wrong_grant_and_later_receipts_block_confirmation(self):
        later = self.receipt(received='2026-05-01')
        other = FinancialService(IDENTITIES['harbour'], self.repo, self.storage)
        for entry in (self.draft(conversionMode='receipt'), self.draft(conversionMode='receipt', receiptId=later['id'])):
            self.assertIsNone(entry['reportAmountMinor'])
            self.deny(lambda: self.confirm(entry))
        self.deny(lambda: other.add_receipt({'grantId': self.grant['id'], 'sourceCurrency': 'RON', 'receivedDate': '2026-04-01', 'sourceAmount': '10', 'rate': '0.3'}), 'not_found')

    def test_exact_rounding_and_rates_are_not_binary_float_math(self):
        self.assertEqual(convert_minor(1, '0.5'), 1)
        self.assertEqual(convert_minor(-1, '0.5'), -1)
        self.assertEqual(normalize_rate('0.321549266958621'), '0.321549266958621')
        for value in ('0', '-1', 'NaN', 'Infinity', '1e3', '1,25'):
            self.deny(lambda: normalize_rate(value))

    def test_confirmed_entries_immutable_adjustment_preserves_rate_and_trace(self):
        original = self.confirm(self.draft())
        self.deny(lambda: self.service.update_entry(original['id'], {'expectedVersion': original['version'], 'amount': '1'}), 'already_handled')
        adjustment = self.service.adjust_entry(original['id'], {'expectedVersion': original['version'], 'amount': '-20', 'reason': 'Fictional eligibility correction'})
        self.confirm(adjustment)
        self.assertEqual(self.service.financial_report(self.grant['id'])['totalMinor'], 2400)
        self.assertEqual(calculate(self.repo.read('brightpath'), self.grant['id'])['allocatedMinor'], 2400)
        stored = self.repo.read('brightpath')['financeEntries'][original['id']]
        self.assertEqual(stored['sourceAmountMinor'], 10000)
        excessive = self.service.adjust_entry(original['id'], {'expectedVersion': original['version'], 'amount': '-100', 'reason': 'Too large'})
        self.deny(lambda: self.confirm(excessive))

    def test_stale_and_invalid_batch_leaves_all_records_unchanged(self):
        good = self.draft()
        bad = self.draft(reference='DEMO-BAD', conversionMode='receipt')
        before = copy.deepcopy(self.repo.read('brightpath'))
        self.deny(lambda: self.service.confirm_entries({'entries': [{'id': e['id'], 'expectedVersion': e['version']} for e in (good, bad)]}))
        self.assertEqual(before, self.repo.read('brightpath'))
        updated = self.service.update_entry(good['id'], {'expectedVersion': good['version'], 'amount': '110'})
        self.deny(lambda: self.confirm(good), 'stale')
        self.confirm(updated)

    def test_foreign_exchange_totals_stay_in_separate_currencies(self):
        self.confirm(self.draft())
        portfolio = self.service.portfolio()
        buckets = {b['currency']: b for b in portfolio['totals']['byCurrency']}
        self.assertIsNone(portfolio['totals']['currency'])
        self.assertEqual(buckets['CAD']['allocatedMinor'], 3000)
        self.assertEqual(buckets['RON']['expenseMinor'], 10000)
        self.assertEqual(buckets['EUR']['allocatedMinor'], 420000)
        self.assertEqual(self.service.grant_detail(self.grant['id'])['grant']['allocatedMinor'], 3000)

    def test_refund_and_transfer_do_not_become_reported_expenses(self):
        self.confirm(self.draft(kind='refund', reference='REFUND-DEMO'))
        self.confirm(self.draft(kind='transfer', reference='TRANSFER-DEMO'))
        self.assertEqual(self.service.financial_report(self.grant['id'])['totalMinor'], 0)
        self.assertEqual(len(self.repo.read('brightpath')['expenses']), 5)

    def test_duplicate_expense_different_upload_does_not_count_twice(self):
        self.confirm(self.draft())
        duplicate = self.draft()
        self.deny(lambda: self.confirm(duplicate), 'duplicate')
        self.assertEqual(self.service.financial_report(self.grant['id'])['totalMinor'], 3000)

    def test_report_period_budgets_and_core_pdf_use_reporting_currency(self):
        self.confirm(self.draft())
        report = self.service.save_report_settings(self.grant['id'], {'periodStart': '2026-04-01', 'periodEnd': '2026-06-30', 'budgets': {'1.4 Goods and Supplies': '50'}})
        self.assertEqual(report['rows'][0], {'category': '1.4 Goods and Supplies', 'budgetMinor': 5000, 'actualMinor': 3000, 'varianceMinor': 2000})
        report = self.service.save_report_settings(self.grant['id'], {'periodStart': '2026-07-01', 'periodEnd': '2026-09-30', 'budgets': {}})
        self.assertEqual(report['totalMinor'], 0)
        old_style = self.service.assemble_report_draft(self.grant['id'])
        self.assertIn('CAD 30.00', old_style['narrative'])
        self.assertEqual(old_style['expenses'][0]['amountMinor'], 3000)
        from grantthread.reports import pdf_bytes
        from pypdf import PdfReader
        text = PdfReader(io.BytesIO(pdf_bytes(old_style))).pages[0].extract_text()
        self.assertIn('CAD 30.00', text)
        self.assertNotIn('EUR', text)

    def test_core_report_records_grant_wide_receipts_in_reporting_currency(self):
        receipt = self.receipt(rate='4', rateDirection='source_per_report')
        self.receipt(amount='200', rate='0.5', received='2026-03-01')
        self.receipt(amount='200', rate='0.25', received='2026-07-01')
        self.service.add_receipt({'grantId': 'digital-belonging', 'sourceCurrency': 'EUR', 'sourceAmount': '900',
                                 'receivedDate': '2026-04-01', 'rate': '1'})
        self.confirm(self.draft(conversionMode='receipt', receiptId=receipt['id']))
        period = self.service.save_report_settings(self.grant['id'], {
            'periodStart': '2026-04-01', 'periodEnd': '2026-06-30', 'budgets': {}})
        self.assertEqual(period['receiptsTotalMinor'], 25000)
        report = self.service.assemble_report_draft(self.grant['id'])
        # Core reports are grant-wide, unlike the period-selected Financials view.
        self.assertEqual(report['confirmedReceiptsMinor'], 40000)
        self.assertEqual(report['allocatedMinor'], 2500)
        from grantthread.reports import pdf_bytes
        from pypdf import PdfReader
        text = PdfReader(io.BytesIO(pdf_bytes(report))).pages[0].extract_text()
        self.assertIn('Recorded funding receipts', text)
        self.assertIn('CAD 400.00', text)
        self.assertIn('CAD 25.00', text)
        self.assertNotIn('Not provided', text)
        self.assertNotIn('EUR', text)

    def test_core_report_distinguishes_absent_receipts_and_rounded_zero(self):
        from grantthread.reports import pdf_bytes
        from pypdf import PdfReader
        missing = self.service.assemble_report_draft(self.grant['id'])
        self.assertIsNone(missing['confirmedReceiptsMinor'])
        self.assertIn('Not provided', PdfReader(io.BytesIO(pdf_bytes(missing))).pages[0].extract_text())
        self.receipt(amount='0.01', rate='0.01')
        zero = self.service.assemble_report_draft(self.grant['id'])
        self.assertEqual(zero['confirmedReceiptsMinor'], 0)
        self.assertIs(type(zero['confirmedReceiptsMinor']), int)
        text = PdfReader(io.BytesIO(pdf_bytes(zero))).pages[0].extract_text()
        self.assertIn('CAD 0.00', text)
        self.assertNotIn('Not provided', text)
        legacy = {k: v for k, v in zero.items() if k != 'confirmedReceiptsMinor'}
        self.assertIn('Not provided', PdfReader(io.BytesIO(pdf_bytes(legacy))).pages[0].extract_text())

    def test_core_report_rejects_invalid_receipts_before_reusing_cached_draft(self):
        receipt = self.receipt(rate='0.25')
        self.service.assemble_report_draft(self.grant['id'])
        cases = [('reportCurrency', 'EUR', 'currency_mismatch')]
        cases += [('reportAmountMinor', value, 'invalid_receipt') for value in (-1, True, 25000.0, '25000', None)]
        for field, value, code in cases:
            with self.subTest(field=field, value=value):
                def corrupt(data):
                    data['fundingReceipts'][receipt['id']] = {**receipt, field: value}
                self.repo.mutate('brightpath', corrupt)
                before = self.repo.read('brightpath')
                self.deny(lambda: self.service.assemble_report_draft(self.grant['id']), code)
                self.assertEqual(self.repo.read('brightpath'), before)

    def test_core_report_replaces_legacy_receipt_placeholder_without_mutating_history(self):
        self.receipt(rate='0.25')
        initial = self.service.assemble_report_draft(self.grant['id'])
        def legacy_placeholder(data):
            data['reports'][initial['id']]['confirmedReceiptsMinor'] = None
        self.repo.mutate('brightpath', legacy_placeholder)
        original = copy.deepcopy(self.repo.read('brightpath')['reports'][initial['id']])
        current = self.service.assemble_report_draft(self.grant['id'])
        self.assertNotEqual(current['id'], original['id'])
        self.assertEqual(current['version'], original['version'] + 1)
        self.assertEqual(current['inputVersion'], original['inputVersion'])
        self.assertEqual(current['confirmedReceiptsMinor'], 25000)
        self.assertEqual(self.repo.read('brightpath')['reports'][original['id']], original)
        self.assertEqual(self.service.assemble_report_draft(self.grant['id'])['id'], current['id'])

    def test_unconfirmed_financial_change_stales_prepared_reports(self):
        report = self.service.assemble_report_draft('digital-belonging')
        self.draft(grantId='digital-belonging', currency='EUR', conversionMode='same_currency')
        self.assertEqual(self.service.report(report['id'])['status'], 'stale')
        self.assertIn('Review draft financial entries', self.service.grant_detail('digital-belonging')['grant']['readiness']['missing'])

    def test_xlsx_import_preview_private_source_draft_confirm_and_exports(self):
        raw = self.workbook()
        item = self.service.upload_import({'name': 'fictional-ledger.xlsx', 'kind': 'ledger', 'grantId': self.grant['id'],
            'contentBase64': base64.b64encode(raw).decode(), 'defaultCurrency': 'RON', 'sheet': 'Sheet', 'headerRow': 1,
            'mapping': {'date': 'A', 'description': 'B', 'reference': 'C', 'amount': 'D', 'currency': 'E', 'category': 'F', 'rate': 'G'}})
        self.assertEqual(item['status'], 'preview')
        self.assertNotIn('objectKey', item)
        self.assertEqual(self.service.source_bytes(item['id'])[0], raw)
        again = self.service.upload_import({'name': 'copy.xlsx', 'kind': 'ledger', 'grantId': self.grant['id'], 'contentBase64': base64.b64encode(raw).decode()})
        self.assertTrue(again['alreadyImported'])
        committed = self.service.commit_import(item['id'], {'expectedVersion': item['version'], 'grantId': self.grant['id']})
        self.assertEqual(self.service.financial_report(self.grant['id'])['totalMinor'], 0)
        self.confirm(self.service.overview()['entries'][0])
        self.assertEqual(self.service.financial_report(self.grant['id'])['totalMinor'], 3215)
        self.deny(lambda: self.service.commit_import(item['id'], {'expectedVersion': committed['version']}))
        other = FinancialService(IDENTITIES['harbour'], self.repo, self.storage)
        self.deny(lambda: other.source_bytes(item['id']), 'not_found')
        from openpyxl import load_workbook
        exported = load_workbook(io.BytesIO(self.service.export_report(self.grant['id'])))
        self.assertTrue(exported.sheetnames)

    def bank_import(self, direction='debit', reference='DEMO-INV-1', **result):
        parsed = {'recognized': True, 'rows': [{'date': '2026-04-20', 'description': 'Fictional bank payment',
                  'reference': reference, 'amount': '100', 'currency': 'RON', 'direction': direction, 'source': {'page': 1}}],
                  'textPages': [{'page': 1, 'text': 'Fictional test source'}], 'warnings': [], **result}
        with patch('grantthread.finance_io.parse_bank_pdf', return_value=parsed):
            return self.service.upload_import({'name': 'fictional.pdf', 'kind': 'bank', 'grantId': self.grant['id'],
                'contentBase64': base64.b64encode((direction + reference).encode()).decode()})

    def test_repeated_small_adjustments_finish_at_exact_zero(self):
        original = self.confirm(self.draft(amount='0.03', manualRate='0.5'))
        deltas = []
        for _ in range(3):
            entry = self.service.adjust_entry(original['id'], {'expectedVersion': original['version'], 'amount': '-0.01', 'reason': 'Fictional correction'})
            deltas.append(self.confirm(entry)['reportAmountMinor'])
        self.assertEqual(deltas, [-1, 0, -1])
        self.assertEqual(self.service.financial_report(self.grant['id'])['totalMinor'], 0)
        self.assertEqual(calculate(self.repo.read('brightpath'), self.grant['id'])['allocatedMinor'], 0)

    def test_confirming_adjustment_invalidates_other_previews_for_same_expense(self):
        original = self.confirm(self.draft(amount='0.03', manualRate='0.5'))
        adjustments = [self.service.adjust_entry(original['id'], {'expectedVersion': original['version'],
            'amount': '-0.01', 'reason': 'Fictional correction'}) for _ in range(2)]
        self.assertEqual([entry['reportAmountMinor'] for entry in adjustments], [-1, -1])
        self.confirm(adjustments[0])
        self.deny(lambda: self.confirm(adjustments[1]), 'stale')
        refreshed = next(entry for entry in self.service.overview()['entries'] if entry['id'] == adjustments[1]['id'])
        self.assertEqual(refreshed['reportAmountMinor'], 0)
        self.assertEqual(self.confirm(refreshed)['reportAmountMinor'], 0)

    def test_batch_cannot_silently_reprice_two_adjustments_of_same_expense(self):
        original = self.confirm(self.draft(amount='0.03', manualRate='0.5'))
        adjustments = [self.service.adjust_entry(original['id'], {'expectedVersion': original['version'],
            'amount': '-0.01', 'reason': 'Fictional correction'}) for _ in range(2)]
        before = copy.deepcopy(self.repo.read('brightpath'))
        self.deny(lambda: self.service.confirm_entries({'entries': [
            {'id': entry['id'], 'expectedVersion': entry['version']} for entry in adjustments]}), 'conflict')
        self.assertEqual(self.repo.read('brightpath'), before)

    def test_independent_adjustments_can_still_be_confirmed_together(self):
        originals = [self.confirm(self.draft(amount='0.03', manualRate='0.5', reference=f'INDEPENDENT-{index}')) for index in range(2)]
        adjustments = [self.service.adjust_entry(original['id'], {'expectedVersion': original['version'],
            'amount': '-0.01', 'reason': 'Fictional independent correction'}) for original in originals]
        confirmed = self.service.confirm_entries({'entries': [
            {'id': entry['id'], 'expectedVersion': entry['version']} for entry in adjustments]})
        self.assertEqual([entry['reportAmountMinor'] for entry in confirmed['entries']], [-1, -1])
        self.assertEqual(self.service.financial_report(self.grant['id'])['totalMinor'], 2)

    def test_adjustment_cannot_exceed_supported_native_amount(self):
        original = self.confirm(self.draft(amount='9999999999.99', manualRate='0.000001'))
        adjustment = self.service.adjust_entry(original['id'], {'expectedVersion': original['version'],
            'amount': '0.01', 'reason': 'Fictional amount outside supported bounds'})
        self.assertIsNone(adjustment['reportAmountMinor'])
        self.deny(lambda: self.confirm(adjustment))

    def test_bank_match_preserves_existing_fx_and_does_not_duplicate_expense(self):
        original = self.confirm(self.draft())
        item = self.bank_import()
        imported = self.service.commit_import(item['id'], {'expectedVersion': item['version']})
        entry = next(e for e in self.service.overview()['entries'] if e['id'] == imported['entryIds'][0])
        self.assertEqual(entry['kind'], 'payment_match')
        changed = self.service.update_entry(entry['id'], {'expectedVersion': entry['version'], 'kind': 'expense', 'conversionMode': 'manual', 'manualRate': '0.3'})
        self.deny(lambda: self.confirm(changed), 'duplicate')
        match = self.service.update_entry(changed['id'], {'expectedVersion': changed['version'], 'kind': 'payment_match'})
        paid = self.confirm(match)
        self.assertEqual(paid['fxRate'], original['fxRate'])
        self.assertEqual(self.service.financial_report(self.grant['id'])['totalMinor'], 3000)
        self.assertEqual(self.repo.read('brightpath')['expenses'][original['expenseId']]['paymentEntryId'], paid['id'])

    def test_credit_needs_classification_and_cannot_turn_into_positive_expense(self):
        item = self.bank_import(direction='credit')
        imported = self.service.commit_import(item['id'], {'expectedVersion': item['version']})
        entry = next(e for e in self.service.overview()['entries'] if e['id'] == imported['entryIds'][0])
        self.assertEqual(entry['kind'], 'unclassified')
        self.deny(lambda: self.confirm(entry))
        entry = self.service.update_entry(entry['id'], {'expectedVersion': entry['version'], 'kind': 'expense',
               'category': '1.4 Goods and Supplies', 'conversionMode': 'manual', 'manualRate': '0.3'})
        self.deny(lambda: self.confirm(entry))
        self.assertEqual(self.service.financial_report(self.grant['id'])['totalMinor'], 0)

    def test_ledger_credit_direction_is_preserved_through_review(self):
        from openpyxl import Workbook
        book = Workbook()
        book.active.append(['Date', 'Description', 'Reference', 'Amount', 'Direction', 'Currency', 'Category', 'Rate'])
        book.active.append(['2026-04-20', 'Fictional refund', 'LEDGER-CREDIT-1', 100, 'credit', 'RON', '1.4 Goods and Supplies', .3])
        raw = io.BytesIO()
        book.save(raw)
        item = self.service.upload_import({'name': 'credit-ledger.xlsx', 'kind': 'ledger', 'grantId': self.grant['id'],
            'contentBase64': base64.b64encode(raw.getvalue()).decode()})
        self.assertEqual(item['rows'][0]['amount'], '-100.00')
        imported = self.service.commit_import(item['id'], {'expectedVersion': item['version']})
        entry = next(entry for entry in self.service.overview()['entries'] if entry['id'] == imported['entryIds'][0])
        self.assertEqual(entry['kind'], 'unclassified')
        changed = self.service.update_entry(entry['id'], {'expectedVersion': entry['version'], 'kind': 'expense'})
        self.deny(lambda: self.confirm(changed))
        refund = self.service.update_entry(changed['id'], {'expectedVersion': changed['version'], 'kind': 'refund'})
        self.confirm(refund)
        self.assertEqual(self.service.financial_report(self.grant['id'])['totalMinor'], 0)

    def test_signed_ledger_refund_cannot_lose_source_direction_during_edit(self):
        from openpyxl import Workbook
        book = Workbook()
        book.active.append(['Date', 'Description', 'Reference', 'Amount', 'Currency', 'Category', 'Rate'])
        book.active.append(['2026-04-20', 'Fictional negative invoice', 'SIGNED-REFUND-1', -20, 'RON', '1.4 Goods and Supplies', .3])
        raw = io.BytesIO()
        book.save(raw)
        item = self.service.upload_import({'name': 'signed-refund.xlsx', 'kind': 'ledger', 'grantId': self.grant['id'],
            'contentBase64': base64.b64encode(raw.getvalue()).decode()})
        imported = self.service.commit_import(item['id'], {'expectedVersion': item['version']})
        changed = self.service.update_entry(imported['entryIds'][0], {'expectedVersion': 1,
            'kind': 'expense', 'sourceDirection': 'debit', 'amount': '25'})
        self.assertEqual(changed['sourceDirection'], 'credit')
        self.deny(lambda: self.confirm(changed))
        self.assertEqual(self.service.financial_report(self.grant['id'])['totalMinor'], 0)

    def test_legacy_imported_credit_drafts_use_immutable_source_direction(self):
        from openpyxl import Workbook
        for explicit_direction in (False, True):
            with self.subTest(explicit_direction=explicit_direction):
                book = Workbook()
                book.active.append(['Date', 'Description', 'Reference', 'Amount', 'Currency', 'Category', 'Rate']
                                   + (['Direction'] if explicit_direction else []))
                book.active.append(['2026-04-20', 'Fictional legacy credit', f'LEGACY-CREDIT-{explicit_direction}',
                    100 if explicit_direction else -100, 'RON', '1.4 Goods and Supplies', .3]
                    + (['credit'] if explicit_direction else []))
                raw = io.BytesIO()
                book.save(raw)
                item = self.service.upload_import({'name': 'legacy-credit.xlsx', 'kind': 'ledger', 'grantId': self.grant['id'],
                    'contentBase64': base64.b64encode(raw.getvalue()).decode()})
                imported = self.service.commit_import(item['id'], {'expectedVersion': item['version']})
                entry_id = imported['entryIds'][0]
                def legacy_record(data):
                    entry = data['financeEntries'][entry_id]
                    entry.pop('sourceDirection')
                    entry['kind'] = 'expense'
                    if explicit_direction:
                        # The previous parser discarded Direction and left a positive cached row.
                        data['financeImports'][item['id']]['rows'][0]['amount'] = '100.00'
                self.repo.mutate('brightpath', legacy_record)
                entry = next(entry for entry in self.service.overview()['entries'] if entry['id'] == entry_id)
                self.deny(lambda: self.confirm(entry))
                self.assertEqual(self.service.financial_report(self.grant['id'])['totalMinor'], 0)
                if explicit_direction:
                    source = self.repo.read('brightpath')['financeImports'][item['id']]
                    self.storage.put(source['objectKey'], b'Fictional corrupted original', source['contentType'])
                    self.deny(lambda: self.confirm(entry), 'source_unavailable')

    def test_legacy_debit_draft_keeps_reviewed_edits_and_can_be_confirmed(self):
        raw = self.workbook()
        item = self.service.upload_import({'name': 'legacy-debit.xlsx', 'kind': 'ledger', 'grantId': self.grant['id'],
            'contentBase64': base64.b64encode(raw).decode()})
        imported = self.service.commit_import(item['id'], {'expectedVersion': item['version']})
        entry_id = imported['entryIds'][0]
        self.repo.mutate('brightpath', lambda data: data['financeEntries'][entry_id].pop('sourceDirection'))
        edited = self.service.update_entry(entry_id, {'expectedVersion': 1, 'amount': '110'})
        self.assertEqual(self.confirm(edited)['sourceAmountMinor'], 11000)

    def test_ledger_upload_honours_and_preserves_explicit_decimal_separator(self):
        from openpyxl import Workbook
        book = Workbook()
        book.active.append(['Date', 'Description', 'Amount', 'Currency', 'Category', 'Rate'])
        book.active.append(['2026-04-20', 'Fictional grouped amount', '1.234', 'RON', '1.4 Goods and Supplies', .3])
        raw = io.BytesIO()
        book.save(raw)
        item = self.service.upload_import({'name': 'grouped-ledger.xlsx', 'kind': 'ledger', 'grantId': self.grant['id'],
            'contentBase64': base64.b64encode(raw.getvalue()).decode(), 'decimalSeparator': ','})
        self.assertEqual(item['status'], 'preview')
        self.assertEqual(item['rows'][0]['amount'], '1234.00')
        self.assertEqual(item['decimalSeparator'], ',')

    def test_one_template_rate_requires_one_source_currency_pair(self):
        self.confirm(self.draft(reference='PAIR-RON'))
        self.confirm(self.draft(currency='USD', reference='PAIR-USD'))
        report = self.service.financial_report(self.grant['id'])
        self.assertIsNone(report['rate'])

    def test_overlapping_bank_sources_cannot_reclassify_an_existing_payment(self):
        original = self.confirm(self.draft())
        item = self.bank_import(reference='BANK-MOVEMENT-7')
        imported = self.service.commit_import(item['id'], {'expectedVersion': item['version']})
        entry = self.service.update_entry(imported['entryIds'][0], {'expectedVersion': 1, 'kind': 'payment_match', 'matchExpenseId': original['expenseId']})
        self.confirm(entry)
        # A different file hash may contain the very same bank movement.
        def overlap(data):
            source = copy.deepcopy(data['financeImports'][item['id']])
            source.update(id='overlap', sha256='different-file', entryIds=[], status='preview', version=1)
            data['financeImports']['overlap'] = source
        self.repo.mutate('brightpath', overlap)
        imported = self.service.commit_import('overlap', {'expectedVersion': 1})
        duplicate = self.service.update_entry(imported['entryIds'][0], {'expectedVersion': 1, 'kind': 'expense',
            'category': '1.4 Goods and Supplies', 'conversionMode': 'manual', 'manualRate': '0.3'})
        self.deny(lambda: self.confirm(duplicate), 'duplicate')
        self.assertEqual(self.service.financial_report(self.grant['id'])['totalMinor'], 3000)

    def test_incomplete_source_stays_visible_until_reviewed_and_keeps_warning(self):
        item = self.bank_import(extractionComplete=False, balanceChecks=[{'matches': False}], warnings=['Fictional missing movement'])
        report = self.service.financial_report(self.grant['id'])
        self.assertEqual(report['unresolvedImportCount'], 1)
        imported = self.service.commit_import(item['id'], {'expectedVersion': item['version']})
        self.deny(lambda: self.service.review_import(item['id'], {'expectedVersion': imported['version'], 'note': 'Checked'}))
        entry = next(e for e in self.service.overview()['entries'] if e['id'] == imported['entryIds'][0])
        self.service.reject_entry(entry['id'], {'expectedVersion': entry['version']})
        self.assertEqual(self.service.financial_report(self.grant['id'])['unresolvedImportCount'], 1)
        self.service.review_import(item['id'], {'expectedVersion': imported['version'], 'note': 'Reviewed statement and entered the missing movements separately.'})
        report = self.service.financial_report(self.grant['id'])
        self.assertEqual(report['unresolvedImportCount'], 0)
        self.assertTrue(report['sourceWarnings'])

    def test_manual_financial_report_is_not_labelled_synthetic(self):
        self.confirm(self.draft())
        self.assertFalse(self.service.financial_report(self.grant['id'])['synthetic'])

    def test_legacy_csv_cannot_reallocate_immutable_financial_expense(self):
        from grantthread.domain import parse_csv
        entry = self.confirm(self.draft(grantId='digital-belonging', currency='EUR', conversionMode='same_currency'))
        csv = 'expense_id,description,amount,currency,date,grant_id,allocation_amount\n' + f"{entry['expenseId']},Fictional workshop supplies,100,EUR,2026-04-15,community-makers,100\n"
        result = parse_csv(csv, self.repo.read('brightpath'))
        self.assertTrue(result['errors'])
        self.assertIn('Financials', str(result['errors']))

    def test_malformed_financial_choices_are_domain_errors(self):
        self.deny(lambda: self.draft(kind=[]))
        self.deny(lambda: self.draft(conversionMode={}))
        self.deny(lambda: self.receipt(rateDirection=[]))
        self.deny(lambda: self.draft(conversionMode='receipt', receiptId=[]))
        self.deny(lambda: self.draft(kind='payment_match', matchExpenseId={}))

    def test_funder_cannot_access_financial_routes(self):
        for method, path, body in [('GET', '/financials', {}), ('POST', '/financials/grants', {}), ('GET', '/financials/report/' + self.grant['id'], {})]:
            self.deny(lambda: dispatch(method, path, body, IDENTITIES['northstar'], self.repo, self.storage), 'forbidden')


if __name__ == '__main__':
    unittest.main()
