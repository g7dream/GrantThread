"""Fictional financial file fixtures only; no real account data or workbooks."""
import io
import re
import sys
import unittest
import zipfile
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from grantthread.errors import DomainError
from grantthread.finance_io import (GCC_CATEGORIES, extract_pdf_pages, export_financial_xlsx,
                                   fill_report_template, inspect_xlsx, parse_bank_csv,
                                   parse_bank_pdf, parse_ledger_xlsx)


def workbook(rows=None, template=False):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "General Ledger"
    headers = [None, "Transaction Date (MM/DD/YYYY)", "Account Name", "Document Number", "Description of Transaction",
               "Expense (RON) ", "Exchange Rate from Local Currency to CAD", "Balance CAD", "Budget Classification", "Funder ID"]
    for col, value in enumerate(headers, 1):
        sheet.cell(4, col, value)
    for number, values in enumerate(rows or [[None, datetime(2026, 1, 12), "Fictional account", "DEMO-01", "Workshop supplies", 123.45, 0.31, 38.27, GCC_CATEGORIES[3], "DEMO"]], 5):
        for col, value in enumerate(values, 1):
            sheet.cell(number, col, value)
    if template:
        report = book.create_sheet("Financial Report")
        report["B2"] = "Fictional report template"
        report["B14"] = "Original period"
        report["E14"] = "Original end"
        report["E16"] = 0.31
        for row, category in enumerate(GCC_CATEGORIES, 25):
            report.cell(row, 2, category)
            report.cell(row, 7, 99)
            report.cell(row, 7).fill = PatternFill("solid", fgColor="AABBCC")
            report.cell(row, 8, f"=G{row}*2")
        report["G33"] = "=SUM(G25:G32)"
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def rewrite_zip(raw, mutate):
    source, output = zipfile.ZipFile(io.BytesIO(raw)), io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            target.writestr(entry.filename, mutate(entry.filename, source.read(entry)))
    return output.getvalue()


def bank_pdf(opening=100, closing=130, unknown=False, no_reference=False):
    from reportlab.pdfgen import canvas
    data = io.BytesIO()
    page = canvas.Canvas(data, pagesize=(700, 800))
    page.setFont("Courier", 10)
    if unknown:
        page.drawString(30, 760, "Fictional unfamiliar statement layout")
    else:
        page.drawString(30, 760, "BANCA TRANSILVANIA - FICTIONAL TEST ONLY")
        page.drawString(30, 740, "EXTRAS CONT - 12/01/2026")
        page.drawString(30, 720, "Valuta RON")
        for x, label in [(30, "Data"), (125, "Descriere"), (490, "Debit"), (620, "Credit")]:
            page.drawString(x, 700, label)
        page.drawString(125, 680, "SOLD ANTERIOR")
        page.drawRightString(655, 680, f"{opening:.2f}")
        page.drawString(30, 660, "12/01/2026")
        page.drawString(125, 660, "Transfer intern - canal electronic")
        page.drawRightString(655, 660, "50.00")
        page.drawString(125, 645, "DEMO-ADJ-2026-03 fictional refund")
        if not no_reference:
            page.drawString(125, 630, "REF: FICTIONAL-CREDIT-01")
        page.drawString(30, 610, "12/01/2026")
        page.drawString(125, 610, "Fictional supplies payment")
        page.drawRightString(515, 610, "20.00")
        if not no_reference:
            page.drawString(125, 595, "REF: FICTIONAL-DEBIT-01")
        page.drawString(30, 575, "12/01/2026 RULAJ ZI")
        page.drawRightString(515, 575, "20.00")
        page.drawRightString(655, 575, "50.00")
        page.drawString(125, 555, "SOLD FINAL CONT")
        page.drawRightString(655, 555, f"{closing:.2f}")
        page.drawString(30, 520, "Acest extras - fictional test only")
    page.save()
    return data.getvalue()


class LedgerImport(unittest.TestCase):
    def test_known_headers_and_template_inspection(self):
        result = inspect_xlsx(workbook(template=True))
        self.assertEqual((result["defaultSheet"], result["headerRow"]), ("General Ledger", 4))
        self.assertEqual(result["mapping"], {"date": "B", "reference": "D", "description": "E", "amount": "F", "rate": "G", "category": "I"})
        self.assertEqual((result["defaultCurrency"], result["suggestedDateFormat"]), ("RON", "mdy"))
        self.assertEqual(result["templateMapping"]["categories"][GCC_CATEGORIES[0]], "G25")

    def test_exact_amount_currency_rate_date_and_source(self):
        result = parse_ledger_xlsx(workbook(), {})
        row = result["rows"][0]
        self.assertEqual((row["amount"], row["rate"], row["currency"], row["date"]), ("123.45", "0.31", "RON", "2026-01-12"))
        self.assertEqual(row["source"], {"sheet": "General Ledger", "row": 5})

    def test_ledger_preserves_high_precision_rates_and_rejects_unsupported_rates(self):
        for value, expected in [(0.321549266958621, "0.321549266958621"),
                                ("0.321549266958621123", "0.321549266958621123")]:
            raw = workbook([[None, datetime(2026, 1, 12), None, "DEMO", "Supplies", 100, value]])
            self.assertEqual(parse_ledger_xlsx(raw, {})["rows"][0]["rate"], expected)
        for value in ["0.3215492669586211234", "0.0000001", "100001"]:
            raw = workbook([[None, datetime(2026, 1, 12), None, "DEMO", "Supplies", 100, value]])
            with self.subTest(value=value), self.assertRaises(DomainError):
                parse_ledger_xlsx(raw, {})

    def test_numeric_cells_ignore_text_locale_and_grouping_is_validated(self):
        result = parse_ledger_xlsx(workbook(), {"decimalSeparator": ","})
        self.assertEqual((result["rows"][0]["amount"], result["rows"][0]["rate"]), ("123.45", "0.31"))
        raw = workbook([[None, datetime(2026, 1, 12), None, "DEMO", "Supplies", "12,34.56", 0.3]])
        with self.assertRaises(DomainError):
            parse_ledger_xlsx(raw, {"decimalSeparator": "."})

    def test_negative_adjustment_kept_totals_and_audit_summary_skipped(self):
        rows = [[None, None, None, None, "Beginning balance", 500],
                [None, datetime(2026, 1, 12), None, "DEMO-ADJ-01", "Audit adjustment refund", -20, 0.3, None, "Supplies"],
                [None, None, None, None, "Total", "=SUM(F5:F6)"],
                [None, None, None, None, "Audit reconciliation summary"],
                [None, datetime(2026, 1, 13), None, "SUMMARY", "Original and revised amount", 999]]
        result = parse_ledger_xlsx(workbook(rows), {})
        self.assertEqual([row["amount"] for row in result["rows"]], ["-20.00"])
        self.assertIn("4", result["warnings"][0])

    def test_dated_total_description_is_an_expense_but_balances_and_aggregates_are_not(self):
        rows = [[None, datetime(2026, 1, 12), None, "DEMO-01", "Total workshop package", 123, 0.3],
                [None, datetime(2026, 1, 12), None, None, "Opening balance", 50, 0.3],
                [None, None, None, None, "Total", "=SUM(F5:F6)"]]
        result = parse_ledger_xlsx(workbook(rows), {})
        self.assertEqual([(r["description"], r["amount"]) for r in result["rows"]], [("Total workshop package", "123.00")])
        self.assertIn("Skipped 2", result["warnings"][0])

    def test_uncached_formula_rejected_not_zero(self):
        rows = [[None, datetime(2026, 1, 12), None, "DEMO-01", "Supplies", "=1+2", 0.3]]
        with self.assertRaises(DomainError) as caught:
            parse_ledger_xlsx(workbook(rows), {})
        self.assertEqual(caught.exception.code, "formula_cache_missing")

    def test_cached_formula_is_read_without_evaluation(self):
        raw = workbook([[None, datetime(2026, 1, 12), None, "DEMO-01", "Supplies", "=1+2", 0.3]])
        raw = rewrite_zip(raw, lambda name, value: re.sub(rb"<f>1\+2</f><v\s*/>", b"<f>1+2</f><v>3</v>", value) if name == "xl/worksheets/sheet1.xml" else value)
        self.assertEqual(parse_ledger_xlsx(raw, {})["rows"][0]["amount"], "3.00")

    def test_unknown_headers_require_explicit_mapping(self):
        book = openpyxl.Workbook()
        book.active.append(["When?", "What?", "Value?"])
        book.active.append(["12/01/2026", "Fictional expense", 22])
        stream = io.BytesIO(); book.save(stream)
        with self.assertRaises(DomainError):
            parse_ledger_xlsx(stream.getvalue(), {})
        result = parse_ledger_xlsx(stream.getvalue(), {"headerRow": 1, "mapping": {"date": "A", "description": "B", "amount": "C"}, "defaultCurrency": "CAD"})
        self.assertEqual(result["rows"][0]["date"], "2026-01-12")

    def test_excel_epoch_and_explicit_text_date_format(self):
        raw = workbook([[None, 44572, None, "DEMO-01", "Supplies", 12, 0.3]])
        book = openpyxl.load_workbook(io.BytesIO(raw)); book.epoch = CALENDAR_MAC_1904
        stream = io.BytesIO(); book.save(stream)
        self.assertEqual(parse_ledger_xlsx(stream.getvalue(), {})["rows"][0]["date"], "2026-01-12")
        raw = workbook([[None, "01/12/2026", None, "DEMO-01", "Supplies", 12, 0.3]])
        self.assertEqual(parse_ledger_xlsx(raw, {"dateFormat": "mdy"})["rows"][0]["date"], "2026-01-12")
        result = parse_ledger_xlsx(raw, {})
        self.assertEqual((result["dateFormat"], result["defaultCurrency"], result["rows"][0]["date"]), ("mdy", "RON", "2026-01-12"))
        self.assertEqual(parse_ledger_xlsx(raw, {"dateFormat": "dmy"})["rows"][0]["date"], "2026-12-01")
        self.assertEqual(parse_ledger_xlsx(raw, {"suggestedDateFormat": "dmy"})["dateFormat"], "dmy")
        mapping = {**result["mapping"], "currency": "", "direction": None}
        self.assertEqual(parse_ledger_xlsx(raw, {"mapping": mapping})["mapping"], result["mapping"])

    def test_summary_income_table_and_audit_comparisons_are_not_expenses(self):
        book = openpyxl.load_workbook(io.BytesIO(workbook()))
        sheet = book.active
        headers = [c.value for c in sheet[4]]
        table = {6: [None, "Current quarter ending balance", None, None, "Current quarter ending balance", "=1.1+2.2"],
                 7: [None, None, None, None, "Summary by category"],
                 8: [None, None, None, None, "Supplies", "=F5"],
                 9: [None, "Transaction Date (MM/DD/YYYY)", None, "Document Number", "Description", "Gain (RON)", "Exchange Rate", None, "Funder ID"],
                 10: [None, datetime(2026, 1, 12), None, "DEMO-INCOME", "Fictional grant receipt", 999, 0.3],
                 12: headers,
                 13: [None, datetime(2026, 1, 13), None, "DEMO-ADJ", "Fictional expense correction", -20, 0.3, None, GCC_CATEGORIES[3]],
                 14: [None, "Audit adjustment comparison"],
                 15: headers,
                 16: [None, datetime(2026, 1, 13), None, "DEMO-COMPARISON", "Original comparison only", 999, 0.3]}
        for number, row in table.items():
            for col, value in enumerate(row, 1):
                sheet.cell(number, col, value)
        source = io.BytesIO(); book.save(source)
        result = parse_ledger_xlsx(source.getvalue(), {})
        self.assertEqual([(r["source"]["row"], r["amount"]) for r in result["rows"]], [(5, "123.45"), (13, "-20.00")])
        self.assertTrue(any("gain/income table at row 9" in warning for warning in result["warnings"]))
        self.assertTrue(any("matching ledger table at row 12" in warning for warning in result["warnings"]))

    def test_missing_date_or_ambiguous_amount_rejected(self):
        for value in [[None, None, None, "DEMO-01", "Supplies", 12], [None, datetime(2026, 1, 12), None, "DEMO-01", "Supplies", "2,500"]]:
            with self.subTest(value=value), self.assertRaises(DomainError):
                parse_ledger_xlsx(workbook([value]), {})

    def test_limits_macros_external_relationships_and_fake_dimensions(self):
        with self.assertRaises(DomainError):
            inspect_xlsx(b"X" * (2 * 1024 * 1024 + 1))
        book = openpyxl.load_workbook(io.BytesIO(workbook())); book.active["BM5"] = 1
        stream = io.BytesIO(); book.save(stream)
        with self.assertRaises(DomainError):
            inspect_xlsx(stream.getvalue())
        raw = rewrite_zip(workbook(), lambda name, value: value.replace(b"sheet.main+xml", b"sheet.macroEnabled.main+xml") if name == "[Content_Types].xml" else value)
        with self.assertRaises(DomainError):
            inspect_xlsx(raw)
        book = openpyxl.load_workbook(io.BytesIO(workbook())); book.active["A1"].hyperlink = "https://example.invalid/"
        stream = io.BytesIO(); book.save(stream)
        self.assertTrue(inspect_xlsx(stream.getvalue())["mapping"])
        raw = rewrite_zip(stream.getvalue(), lambda name, value: value.replace(b"/hyperlink", b"/externalLink") if name.endswith(".rels") else value)
        with self.assertRaises(DomainError):
            inspect_xlsx(raw)
        book.active["A1"].hyperlink = "file:///example.xlsx"
        stream = io.BytesIO(); book.save(stream)
        with self.assertRaises(DomainError):
            inspect_xlsx(stream.getvalue())
        raw = rewrite_zip(workbook(), lambda name, value: re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', value) if name == "xl/worksheets/sheet1.xml" else value)
        self.assertEqual(len(parse_ledger_xlsx(raw, {})["rows"]), 1)

    def test_zip_bomb_and_row_count_limits(self):
        raw = workbook([[None, datetime(2026, 1, 12), None, "DEMO", "Supplies", 1]] * 501)
        with self.assertRaises(DomainError):
            parse_ledger_xlsx(raw, {})
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("padding.xml", b"0" * 100_000)
        with self.assertRaises(DomainError) as caught:
            inspect_xlsx(output.getvalue())
        self.assertEqual(caught.exception.code, "file_limit")

    def test_bounded_reference_columns_and_formatting_do_not_block_ledger_import(self):
        book = openpyxl.load_workbook(io.BytesIO(workbook(template=True)))
        book["General Ledger"]["BB5"].fill = PatternFill("solid", fgColor="AABBCC")
        reference = book.create_sheet("Fictional reference")
        reference.append([f"Reference {i}" for i in range(53)])
        reference.append(list(range(53)))
        source = io.BytesIO(); book.save(source)
        info = inspect_xlsx(source.getvalue())
        self.assertEqual(len(info["sheets"]), 3)
        self.assertTrue(all(len(row) <= 50 for sheet in info["sheets"] for row in sheet["previewRows"]))
        self.assertEqual(len(parse_ledger_xlsx(source.getvalue(), {})["rows"]), 1)
        with self.assertRaises(DomainError):
            parse_ledger_xlsx(source.getvalue(), {"mapping": {**info["mapping"], "amount": "AY"}})
        report = {"rows": [{"category": GCC_CATEGORIES[0], "actualMinor": 123}]}
        filled = openpyxl.load_workbook(io.BytesIO(fill_report_template(source.getvalue(), report)))
        self.assertEqual(filled["Fictional reference"]["BA2"].value, 52)
        self.assertEqual(filled["General Ledger"]["BB5"].fill.fgColor.rgb, "00AABBCC")


class BankImport(unittest.TestCase):
    def test_csv_explicit_direction_and_split_columns(self):
        raw = b"date,description,debit,credit,currency\n2026-01-12,Refund,,50.00,RON\n2026-01-12,Supplies,20.00,,RON\n"
        rows = parse_bank_csv(raw)["rows"]
        self.assertEqual([(r["direction"], r["amount"]) for r in rows], [("credit", "50.00"), ("debit", "20.00")])
        raw = b"date,description,amount,direction,currency\n2026-01-12,Refund,50.00,credit,RON\n"
        self.assertEqual(parse_bank_csv(raw)["rows"][0]["direction"], "credit")

    def test_csv_ambiguous_direction_dual_amount_and_bad_shape_rejected(self):
        for raw in [b"date,description,amount,currency\n2026-01-12,Payment,50,RON\n",
                    b"date,description,debit,credit,currency\n2026-01-12,Payment,10,20,RON\n",
                    b"date,description,amount,direction,currency\n2026-01-12,Payment,50,debit\n"]:
            with self.subTest(raw=raw), self.assertRaises(DomainError):
                parse_bank_csv(raw)

    def test_pdf_transactions_only_reconcile_and_include_sources(self):
        result = parse_bank_pdf(bank_pdf())
        self.assertTrue(result["recognized"], result)
        self.assertEqual(len(result["rows"]), 2)
        self.assertEqual([(r["direction"], r["amount"]) for r in result["rows"]], [("credit", "50.00"), ("debit", "20.00")])
        self.assertEqual(result["rows"][0]["reference"], "DEMO-ADJ-2026-03")
        self.assertEqual(result["rows"][0]["source"]["page"], 1)
        self.assertFalse(result["warnings"])

    def test_pdf_without_reference_does_not_drop_final_transaction(self):
        result = parse_bank_pdf(bank_pdf(no_reference=True))
        self.assertTrue(result["recognized"])
        self.assertEqual(len(result["rows"]), 2)

    def test_pdf_unknown_and_balance_mismatch_withhold_rows(self):
        for raw in [bank_pdf(unknown=True), bank_pdf(closing=999)]:
            result = parse_bank_pdf(raw)
            self.assertFalse(result["recognized"])
            self.assertEqual(result["rows"], [])
            self.assertTrue(result["textPages"])

    def test_invalid_and_image_only_pdf_rejected(self):
        with self.assertRaises(DomainError):
            extract_pdf_pages(b"not pdf")
        from reportlab.pdfgen import canvas
        output = io.BytesIO(); page = canvas.Canvas(output); page.showPage(); page.save()
        with self.assertRaises(DomainError) as caught:
            extract_pdf_pages(output.getvalue())
        self.assertEqual(caught.exception.code, "ocr_unsupported")


class FinancialExports(unittest.TestCase):
    def report(self):
        return {"synthetic": False, "currency": "CAD", "periodStart": "2026-01-01", "periodEnd": "2026-01-31", "rate": "0.31",
                "rows": [{"category": category, "budgetMinor": 10000, "actualMinor": -123 if index == 0 else index * 100, "varianceMinor": 10000 - (-123 if index == 0 else index * 100)} for index, category in enumerate(GCC_CATEGORIES)]}

    def test_export_exact_negative_numbers_and_formula_safe_strings(self):
        entries = [{"date": "2026-01-12", "description": '=HYPERLINK("https://example.invalid")', "reference": "+TEST", "category": "Supplies",
                    "sourceAmountMinor": -12345, "currency": "RON", "reportAmountMinor": -3827, "reportCurrency": "CAD", "fxRate": "0.31", "source": {"sheet": "Ledger", "row": 5}, "history": [{"action": "adjustment"}]}]
        receipts = [{"receivedDate": "2026-01-12", "description": "=unsafe", "reference": "@name", "sourceAmountMinor": 50000, "sourceCurrency": "RON", "reportAmountMinor": 15500, "reportCurrency": "CAD", "rate": "0.31"}]
        raw = export_financial_xlsx(self.report(), entries, receipts)
        book = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
        self.assertEqual(book["Ledger"]["E4"].value, -123.45)
        self.assertEqual(book["Ledger"]["G4"].value, -38.27)
        self.assertEqual(book["Ledger"]["I4"].value, "0.31")
        self.assertEqual(book["Ledger"]["I4"].data_type, "s")
        self.assertEqual(book["Ledger"]["B4"].data_type, "s")
        self.assertEqual(book["Funding Receipts"]["B4"].data_type, "s")
        self.assertEqual(book["Funding Receipts"]["F4"].value, 155)
        self.assertEqual(book["Financial Report"]["A1"].value, "Financial report")

    def test_export_preserves_exact_rates_receipt_context_and_report_metadata(self):
        report = {**self.report(), "grant": {"name": "Fictional workshop grant"}, "unresolvedCount": 2,
                  "unresolvedImportCount": 1, "sourceWarnings": ["Fictional source needs review"]}
        rate = "0.321549266958621123"
        entries = [{"sourceAmountMinor": 10000, "reportAmountMinor": 3215, "currency": "RON", "fxRate": rate}]
        receipts = [{"sourceAmountMinor": 10000, "sourceCurrency": "RON", "reportAmountMinor": 3215, "rate": rate,
                     "note": "=literal reference", "enteredRate": "3.109943319487937442", "enteredDirection": "source_per_report"}]
        book = openpyxl.load_workbook(io.BytesIO(export_financial_xlsx(report, entries, receipts)))
        self.assertEqual(book["Ledger"]["I4"].value, rate)
        self.assertEqual(book["Funding Receipts"]["H4"].value, rate)
        self.assertEqual(book["Funding Receipts"]["K4"].value, "=literal reference")
        self.assertEqual(book["Funding Receipts"]["K4"].data_type, "s")
        self.assertEqual(book["Funding Receipts"]["L4"].value, "source_per_report")
        self.assertEqual(book["Funding Receipts"]["M4"].value, "3.109943319487937442")
        self.assertEqual(book["Funding Receipts"]["M4"].data_type, "s")
        sheet = book["Financial Report"]
        metadata = {sheet.cell(row, 6).value: sheet.cell(row, 7).value for row in range(3, 10)}
        self.assertEqual(metadata, {"Grant": "Fictional workshop grant", "Reporting currency": "CAD",
            "Period start": "2026-01-01", "Period end": "2026-01-31", "Unresolved entries": 2,
            "Unresolved imports": 1, "Source warnings": 'Fictional source needs review'})
        self.assertEqual((sheet["A12"].value, sheet["C12"].value), ("Total", 26.77))

    def test_template_only_mapped_cells_change_preserves_formulas_styles(self):
        source = workbook(template=True)
        raw = fill_report_template(source, self.report(), {"periodStartCell": "B14", "periodEndCell": "E14", "rateCell": "E16"})
        book = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
        original = openpyxl.load_workbook(io.BytesIO(source), data_only=False)
        self.assertEqual(book["Financial Report"]["G25"].value, -1.23)
        self.assertEqual(book["Financial Report"]["H25"].value, "=G25*2")
        self.assertEqual(book["Financial Report"]["G25"].fill.fgColor.rgb, "00AABBCC")
        self.assertEqual(book["Financial Report"]["B2"].value, original["Financial Report"]["B2"].value)
        self.assertEqual(book["General Ledger"]["F5"].value, original["General Ledger"]["F5"].value)
        self.assertEqual(book["Financial Report"]["B14"].value, "2026-01-01")
        self.assertEqual(book["Financial Report"]["G33"].value, 26.77)

    def test_template_preserves_array_formula_text_and_range(self):
        from openpyxl.worksheet.formula import ArrayFormula
        book = openpyxl.load_workbook(io.BytesIO(workbook(template=True)))
        book["Financial Report"]["J25"] = ArrayFormula(ref="J25:J32", text="=G25:G32*2")
        source = io.BytesIO(); book.save(source)
        filled = openpyxl.load_workbook(io.BytesIO(fill_report_template(source.getvalue(), self.report())))
        formula = filled["Financial Report"]["J25"].value
        self.assertIsInstance(formula, ArrayFormula)
        self.assertEqual((formula.ref, formula.text), ("J25:J32", "=G25:G32*2"))

    def test_template_rate_direction_is_explicit_and_inverse_preserves_precision(self):
        source = workbook(template=True)
        report = {**self.report(), "rate": "0.3"}
        for direction, expected in [("report_per_source", "0.3"), ("source_per_report", "3.333333333333333333")]:
            filled = openpyxl.load_workbook(io.BytesIO(fill_report_template(source, report, {"rateCell": "E16", "rateDirection": direction})))
            self.assertEqual(filled["Financial Report"]["E16"].value, expected)
            self.assertEqual(filled["Financial Report"]["E16"].data_type, "s")
            self.assertEqual(filled["Financial Report"]["H25"].value, "=G25*2")
            self.assertEqual(filled["Financial Report"]["B14"].value, "Original period")
        for direction in ["CAD/RON", "", None, []]:
            with self.subTest(direction=direction), self.assertRaises(DomainError) as caught:
                fill_report_template(source, report, {"rateCell": "E16", "rateDirection": direction})
            self.assertEqual(caught.exception.code, "invalid_mapping")

    def test_template_range_mapping_clears_unused_categories_and_preserves_other_cells(self):
        report = self.report()
        report["rows"] = [report["rows"][0]]
        report["rate"] = "0.321549266958621123"
        mapping = {"sheet": "Financial Report", "categoryColumn": "B", "actualColumn": "G", "startRow": 25,
                   "endRow": 32, "totalCell": "G33", "rateCell": "E16"}
        book = openpyxl.load_workbook(io.BytesIO(fill_report_template(workbook(template=True), report, mapping)))
        sheet = book["Financial Report"]
        self.assertEqual(sheet["G25"].value, -1.23)
        self.assertEqual([sheet.cell(row, 7).value for row in range(26, 33)], [0] * 7)
        self.assertEqual(sheet["G33"].value, -1.23)
        self.assertEqual(sheet["H25"].value, "=G25*2")
        self.assertEqual(sheet["B25"].value, GCC_CATEGORIES[0])
        self.assertEqual(sheet["G25"].fill.fgColor.rgb, "00AABBCC")
        self.assertEqual(sheet["E16"].value, report["rate"])
        self.assertEqual(sheet["E16"].data_type, "s")

    def test_template_custom_range_used_instead_of_default_layout(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Custom report"
        sheet["C6"] = "Fictional category"
        sheet["E6"] = 999
        sheet["F6"] = "=E6*2"
        source = io.BytesIO(); book.save(source)
        report = {"rows": [{"category": "Fictional category", "actualMinor": 1234}]}
        mapping = {"sheet": "Custom report", "categoryColumn": "C", "actualColumn": "E", "startRow": 6, "endRow": 6, "totalCell": "E7"}
        filled = openpyxl.load_workbook(io.BytesIO(fill_report_template(source.getvalue(), report, mapping)))
        self.assertEqual(filled["Custom report"]["E6"].value, 12.34)
        self.assertEqual(filled["Custom report"]["E7"].value, 12.34)
        self.assertEqual(filled["Custom report"]["F6"].value, "=E6*2")

    def test_template_range_rejects_bad_bounds_missing_labels_and_duplicate_categories(self):
        base = {"categoryColumn": "B", "actualColumn": "G", "startRow": 25, "endRow": 32}
        for change in [{"startRow": 24}, {"endRow": 24}, {"endRow": 2001}, {"categoryColumn": "G"}, {"actualColumn": "ZZ"},
                       {"startRow": True}, {"categories": {"Example": "G25"}}]:
            with self.subTest(change=change), self.assertRaises(DomainError):
                fill_report_template(workbook(template=True), self.report(), {**base, **change})
        book = openpyxl.load_workbook(io.BytesIO(workbook(template=True)))
        book["Financial Report"]["B26"] = book["Financial Report"]["B25"].value
        source = io.BytesIO(); book.save(source)
        with self.assertRaises(DomainError):
            fill_report_template(source.getvalue(), self.report(), base)

    def test_template_mapping_missing_category_overlap_and_unknown_layout_rejected(self):
        source = workbook(template=True)
        for mapping in [{"categories": {GCC_CATEGORIES[0]: "G25"}},
                        {"categories": {key: "G25" for key in GCC_CATEGORIES}},
                        {"sheet": "General Ledger"}]:
            with self.subTest(mapping=mapping), self.assertRaises(DomainError):
                fill_report_template(source, self.report(), mapping)


if __name__ == "__main__":
    unittest.main()
