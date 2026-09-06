"""Bounded financial file parsing. Files are data; formulas and macros never run.

This module performs no persistence, model calls or grant allocation. Callers retain
the source file/version and authorise both import confirmation and report exports.
Ledger amounts are signed (credit/refund negative); bank rows have positive amounts
with explicit debit/credit direction. Exchange rates are preserved, not applied here.
"""
import csv
import io
import json
import re
import unicodedata
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, localcontext
from xml.etree import ElementTree

from .errors import DomainError, require


XLSX_MAX_BYTES = 2 * 1024 * 1024
ZIP_MAX_BYTES = 20 * 1024 * 1024
MAX_SHEETS, MAX_ROWS, MAX_COLUMNS, MAX_TRANSACTIONS = 30, 2000, 50, 500
# Reference sheets and formatting may extend beyond the 50 importable columns.
MAX_PHYSICAL_COLUMNS = 64
FIELDS = {"date", "description", "reference", "amount", "debit", "credit", "currency", "category", "rate", "direction"}
GCC_CATEGORIES = ["1.1 Remuneration", "1.2 Subcontractor Fees", "1.3 Travel Costs", "1.4 Goods and Supplies",
                  "1.5 Equipment Costs", "1.6 Project Administration Costs", "1.7 Sub-Grants less Sub-Grantee Indirects", "1.8 Indirect Costs"]
ALIASES = {
    "date": {"date", "transaction date", "transaction date mm dd yyyy", "transaction date dd mm yyyy", "data", "data tranzactiei"},
    "description": {"description", "description of transaction", "transaction description", "descriere", "details"},
    "reference": {"reference", "document number", "reference number", "transaction reference", "referinta", "document no"},
    "amount": {"amount", "expense", "expense ron", "transaction amount", "suma"},
    "debit": {"debit", "debits", "withdrawal", "withdrawals"},
    "credit": {"credit", "credits", "deposit", "deposits"},
    "currency": {"currency", "currency code", "valuta"},
    "category": {"category", "budget classification", "budget category", "categorie"},
    "rate": {"rate", "exchange rate", "fx rate", "exchange rate from local currency to cad"},
    "direction": {"direction", "type", "debit credit"},
}
SUMMARY = re.compile(r"^(?:beginning|opening|closing|ending|added)\s+balance\b|^balance\s+(?:brought|carried|added)\b|^total(?:s|\s|$)|^subtotal\b|^grand total\b|^sold\s+(?:initial|final|anterior)\b|^rulaj\b", re.I)
BALANCE = re.compile(r"^(?:beginning|opening|closing|ending|added)\s+balance\b|^balance\s+(?:brought|carried|added)\b|^sold\s+(?:initial|final|anterior)\b", re.I)
SUMMARY_SECTION = re.compile(r"^audit\s+(?:reconciliation|adjustment)\b|^(?:original|revised)\s+(?:report\s+)?summary|^summary\s+of\s+(?:original|revised)", re.I)


def _normal(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(c for c in text if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def _xlsx(raw, data_only=False, read_only=True):
    require(isinstance(raw, bytes) and 0 < len(raw) <= XLSX_MAX_BYTES, "XLSX must contain 1 byte to 2 MB", "file_limit", 413)
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            require(len(entries) <= 5000, "Workbook contains too many ZIP entries", "file_limit", 413)
            names = [entry.filename for entry in entries]
            require(len(names) == len(set(names)), "Workbook has duplicate ZIP entries", "invalid_workbook", 422)
            total = sum(entry.file_size for entry in entries)
            require(total <= ZIP_MAX_BYTES and total <= max(1, sum(e.compress_size for e in entries)) * 50,
                    "Workbook exceeds expanded size or compression limits", "file_limit", 413)
            for entry in entries:
                name = entry.filename.lower()
                require(not entry.flag_bits & 1, "Encrypted workbook entries are unsupported", "unsupported_workbook", 415)
                require(not name.startswith("/") and ".." not in name.split("/") and "\\" not in name,
                        "Workbook has an invalid ZIP path", "invalid_workbook", 422)
                require("vbaproject" not in name and not name.startswith("xl/externallinks/") and not name.startswith("xl/embeddings/"),
                        "Macros, external links and embedded objects are unsupported", "unsupported_workbook", 415)
                if name.endswith(".xml") or name.endswith(".rels"):
                    content = archive.read(entry)
                    require(b"<!DOCTYPE" not in content.upper() and b"<!ENTITY" not in content.upper(),
                            "XML entity declarations are unsupported", "unsupported_workbook", 415)
                    require(b"macroenabled" not in content.lower(), "Macro-enabled workbooks are unsupported", "unsupported_workbook", 415)
                    if name.endswith(".rels"):
                        xml = ElementTree.fromstring(content)
                        for item in xml:
                            if item.attrib.get("TargetMode", "").lower() == "external":
                                require(item.attrib.get("Type", "").endswith("/hyperlink") and
                                        re.match(r"^https?://", item.attrib.get("Target", ""), re.I),
                                        "External workbook links and non-web hyperlinks are unsupported", "unsupported_workbook", 415)
                    elif name.startswith("xl/worksheets/") and name.endswith(".xml"):
                        xml = ElementTree.fromstring(content)
                        from openpyxl.utils import column_index_from_string
                        physical_rows = [item for item in xml.iter() if item.tag.rsplit("}", 1)[-1] == "row"]
                        require(len(physical_rows) <= MAX_ROWS and all(0 < int(item.attrib.get("r", "0")) <= MAX_ROWS for item in physical_rows),
                                "Sheet exceeds 2,000 physical rows", "file_limit", 413)
                        for item in xml.iter():
                            if item.tag.rsplit("}", 1)[-1] == "c":
                                address = re.fullmatch(r"([A-Z]+)([1-9]\d*)", item.attrib.get("r", ""))
                                require(address and column_index_from_string(address.group(1)) <= MAX_PHYSICAL_COLUMNS and int(address.group(2)) <= MAX_ROWS,
                                        "Sheet exceeds 64 physical columns or 2,000 physical rows", "file_limit", 413)
            require("[Content_Types].xml" in names and "xl/workbook.xml" in names, "File is not an XLSX workbook", "invalid_workbook", 422)
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(raw), data_only=data_only, read_only=read_only, keep_links=False)
        require(0 < len(book.worksheets) <= MAX_SHEETS, "Workbook must have 1 to 30 sheets", "file_limit", 413)
        for sheet in book.worksheets:
            if read_only:
                # Ignore untrusted dimension metadata; discover actual bounded cells.
                sheet.reset_dimensions()
                sheet.calculate_dimension(force=True)
            require(sheet.max_row <= MAX_ROWS and sheet.max_column <= MAX_PHYSICAL_COLUMNS,
                    "Each sheet is limited to 2,000 physical rows and 64 physical columns", "file_limit", 413)
        return book
    except DomainError:
        raise
    except (zipfile.BadZipFile, OSError, ValueError, KeyError, ElementTree.ParseError) as exc:
        raise DomainError("The XLSX workbook could not be read", "invalid_workbook", 422) from exc


def _header(values, complete=True):
    from openpyxl.utils import get_column_letter
    mapping = {}
    for index, value in enumerate(values, 1):
        for field, aliases in ALIASES.items():
            if _normal(value) in aliases:
                if field in mapping:
                    return {}  # A duplicate recognised header is ambiguous.
                mapping[field] = get_column_letter(index)
    return mapping if not complete or {"date", "description"}.issubset(mapping) and ("amount" in mapping or {"debit", "credit"}.issubset(mapping)) else {}


def _preview_value(cell):
    if cell.data_type == "f":
        return "[formula; cached value required for import]"
    value = cell.value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, str):
        return value[:500]
    return value


def inspect_xlsx(raw):
    book = _xlsx(raw)
    try:
        sheets, candidates = [], []
        for sheet in book.worksheets:
            rows = list(sheet.iter_rows(max_row=min(sheet.max_row, 12), max_col=min(sheet.max_column, MAX_COLUMNS)))
            sheets.append({"name": sheet.title, "rowCount": sheet.max_row,
                           "previewRows": [[_preview_value(c) for c in row] for row in rows]})
            for number, row in enumerate(rows, 1):
                mapping = _header([c.value for c in row])
                if mapping:
                    candidates.append((sheet.title, number, mapping, row))
                    break
        chosen = next((c for c in candidates if "general ledger" in c[0].lower()), candidates[0] if candidates else None)
        result = {"sheets": sheets, "defaultSheet": chosen[0] if chosen else book.worksheets[0].title,
                  "headerRow": chosen[1] if chosen else None, "mapping": chosen[2] if chosen else {}}
        if chosen:
            headers = " ".join(str(c.value or "") for c in chosen[3])
            match = re.search(r"(?:Expense|Amount)\s*\(([A-Z]{3})\)", headers, re.I)
            if match:
                result["defaultCurrency"] = match.group(1).upper()
            if "MM/DD/YYYY" in headers.upper():
                result["suggestedDateFormat"] = "mdy"
        if "Financial Report" in book.sheetnames:
            template = book["Financial Report"]
            labels = [str(template.cell(row, 2).value or "").strip() for row in range(25, 33)]
            if [_normal(x) for x in labels] == [_normal(x) for x in GCC_CATEGORIES]:
                result["templateMapping"] = {"sheet": "Financial Report", "categories": {label: f"G{row}" for row, label in enumerate(labels, 25)}, "totalCell": "G33"}
                result["templateHints"] = {"rateCell": "E16", "periodStartCell": "B14", "periodEndCell": "E14"}
        return result
    finally:
        book.close()


def _mapping(value, max_columns, bank=False):
    from openpyxl.utils import column_index_from_string
    require(isinstance(value, dict) and bool(value) and set(value).issubset(FIELDS),
            "Provide an explicit supported column mapping", "mapping_required", 422)
    result = {}
    for key, column in value.items():
        if column in (None, ""):
            continue
        require(isinstance(column, str) and bool(re.fullmatch(r"[A-Za-z]{1,2}", column)), "Map columns using letters such as B or F", "invalid_mapping", 422)
        index = column_index_from_string(column.upper())
        require(index <= min(max_columns, MAX_COLUMNS), "Mapped column is outside the sheet", "invalid_mapping", 422)
        result[key] = index - 1
    require(len(set(result.values())) == len(result), "A column cannot have two field mappings", "invalid_mapping", 422)
    require({"date", "description"}.issubset(result) and ("amount" in result or {"debit", "credit"}.issubset(result)),
            "Map date, description and amount, or both debit and credit", "mapping_required", 422)
    require(not ("amount" in result and ("debit" in result or "credit" in result)), "Choose signed amount or separate debit/credit columns", "invalid_mapping", 422)
    return result


def _number(value, label, places=2, decimal_separator=None):
    require(not isinstance(value, bool) and value is not None and value != "", f"{label} is missing", "invalid_financial_value", 422)
    require(isinstance(value, (str, int, float, Decimal)), f"{label} must be a decimal number", "invalid_financial_value", 422)
    text = str(value).strip()
    if isinstance(value, str):
        text = text.replace("\u00a0", "").replace(" ", "")
        if text.startswith("(") and text.endswith(")"):
            text = "-" + text[1:-1]
        separator = decimal_separator
        if separator is None and "," in text and "." in text:
            separator = "." if text.rfind(".") > text.rfind(",") else ","
        if separator == ",":
            require(bool(re.fullmatch(r"[+-]?(?:\d+|\d{1,3}(?:\.\d{3})+)(?:,\d+)?", text)), f"{label} has invalid digit grouping", "invalid_financial_value", 422)
            text = text.replace(".", "").replace(",", ".")
        elif separator == ".":
            require(bool(re.fullmatch(r"[+-]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?", text)), f"{label} has invalid digit grouping", "invalid_financial_value", 422)
            text = text.replace(",", "")
        elif "," in text:
            require(not re.fullmatch(r"[+-]?\d{1,3}(?:,\d{3})+", text),
                    f"{label} has ambiguous separators; specify decimalSeparator", "ambiguous_amount", 422)
            text = text.replace(",", ".")
        require(bool(re.fullmatch(r"[+-]?\d+(?:\.\d+)?", text)), f"{label} must be a finite decimal number", "invalid_financial_value", 422)
    try:
        number = Decimal(text)
        require(number.is_finite() and abs(number) <= Decimal("9999999999.99"), f"{label} is outside supported bounds", "invalid_financial_value", 422)
        require(number == number.quantize(Decimal(1).scaleb(-places)), f"{label} has too many decimal places", "invalid_financial_value", 422)
        return number
    except InvalidOperation as exc:
        raise DomainError(f"{label} is invalid", "invalid_financial_value", 422) from exc


def _day(value, epoch=None, date_format="dmy"):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (float, int)) and not isinstance(value, bool):
        from openpyxl.utils.datetime import from_excel
        require(epoch is not None, "Numeric dates require an Excel workbook epoch", "invalid_date", 422)
        try:
            converted = from_excel(value, epoch=epoch)
            require(isinstance(converted, datetime), "Excel value does not contain a calendar date", "invalid_date", 422)
            return converted.date().isoformat()
        except (OverflowError, ValueError) as exc:
            raise DomainError("Excel date is invalid", "invalid_date", 422) from exc
    require(isinstance(value, str), "Transaction date is required", "invalid_date", 422)
    text = value.strip()
    patterns = ["%Y-%m-%d"] + ({"dmy": ["%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"],
                                "mdy": ["%m/%d/%Y", "%m-%d-%Y", "%m.%d.%Y"], "ymd": ["%Y/%m/%d"]}[date_format])
    for pattern in patterns:
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            pass
    raise DomainError("Transaction date does not match the selected date format", "invalid_date", 422)


def _rate_text(value, decimal_separator=None):
    rate = _number(value, "Exchange rate", places=18, decimal_separator=decimal_separator)
    require(Decimal("0.000001") <= rate <= Decimal("100000"),
            "Exchange rate must be between 0.000001 and 100000", "invalid_financial_value", 422)
    return format(rate.normalize(), "f")


def _options(options):
    require(isinstance(options, dict), "Import options must be an object")
    require(isinstance(options.get("dateFormat", "dmy"), str) and options.get("dateFormat", "dmy") in {"dmy", "mdy", "ymd"}, "Choose dmy, mdy or ymd date format", "invalid_mapping", 422)
    require(options.get("decimalSeparator") is None or isinstance(options.get("decimalSeparator"), str) and options["decimalSeparator"] in {".", ","}, "Decimal separator must be a dot or comma", "invalid_mapping", 422)


def _transaction(values, columns, options, source, epoch=None, bank=False):
    read = lambda key: values[columns[key]] if key in columns else None
    description = str(read("description") or "").strip()
    require(description, "Transaction description is required", "invalid_financial_value", 422)
    currency = str(read("currency") or options.get("defaultCurrency") or "").strip().upper()
    require(bool(re.fullmatch(r"[A-Z]{3}", currency)), "Choose the source currency explicitly", "currency_required", 422)
    separator = options.get("decimalSeparator")
    if "amount" in columns:
        amount = _number(read("amount"), "Amount", decimal_separator=separator)
        direction = str(read("direction") or "").strip().lower()
        if bank:
            require(direction in {"debit", "credit"}, "Bank amount rows need an explicit debit/credit direction", "mapping_required", 422)
            require(amount >= 0, "Bank amount must be positive; use direction for debit/credit", "invalid_financial_value", 422)
    else:
        debit = _number(read("debit") or 0, "Debit", decimal_separator=separator)
        credit = _number(read("credit") or 0, "Credit", decimal_separator=separator)
        require(debit >= 0 and credit >= 0 and not (debit and credit), "A transaction cannot contain both debit and credit amounts", "invalid_financial_value", 422)
        amount, direction = debit - credit, "debit" if debit else "credit"
        if bank:
            amount = abs(amount)
    require(amount != 0, "A transaction amount cannot be zero", "invalid_financial_value", 422)
    result = {"date": _day(read("date"), epoch, options.get("dateFormat", "dmy")), "description": description,
              "reference": str(read("reference") or "").strip(), "amount": f"{amount:.2f}", "currency": currency,
              "category": str(read("category") or "").strip(), "source": source}
    if bank:
        result["direction"] = direction
    if "rate" in columns and read("rate") not in (None, ""):
        result["rate"] = _rate_text(read("rate"), decimal_separator=separator)
    return result


def parse_ledger_xlsx(raw, options):
    require(isinstance(options, dict), "Import options must be an object")
    info = inspect_xlsx(raw)
    options = {**options, "dateFormat": options.get("dateFormat") or options.get("suggestedDateFormat") or info.get("suggestedDateFormat") or "dmy"}
    _options(options)
    sheet_name = options.get("sheet") or info["defaultSheet"]
    header_row = options.get("headerRow", info["headerRow"])
    mapping = options.get("mapping") or (info["mapping"] if sheet_name == info["defaultSheet"] else {})
    require(type(header_row) is int and 1 <= header_row < MAX_ROWS, "Select the one-based header row", "mapping_required", 422)
    formulas, cached = _xlsx(raw), _xlsx(raw, data_only=True)
    try:
        require(isinstance(sheet_name, str) and sheet_name in cached.sheetnames, "Select an existing ledger sheet", "invalid_mapping", 422)
        sheet, original = cached[sheet_name], formulas[sheet_name]
        columns = _mapping(mapping, sheet.max_column)
        actual_options = {**options}
        if "defaultCurrency" not in actual_options and sheet_name == info["defaultSheet"] and info.get("defaultCurrency"):
            actual_options["defaultCurrency"] = info["defaultCurrency"]
        rows, skipped, section, warnings = [], 0, None, []
        for number, (cells, raw_cells) in enumerate(zip(sheet.iter_rows(min_row=header_row + 1), original.iter_rows(min_row=header_row + 1)), header_row + 1):
            values = [c.value for c in cells]
            labels = [str(c.value or "").strip() for c in raw_cells if c.data_type != "f"]
            if not any(c.value is not None for c in raw_cells):
                continue
            try:
                _day(values[columns["date"]], cached.epoch, actual_options["dateFormat"])
                dated = True
            except DomainError:
                dated = False
            if not dated and any(SUMMARY_SECTION.search(label) for label in labels):
                section = "audit"
            elif section != "audit" and not dated and any(re.match(r"^summary\b", label, re.I) for label in labels):
                section = "summary"
            headers = _header(values, complete=False)
            if section != "audit" and {"date", "description"}.issubset(headers):
                if any(re.match(r"^(?:gain|income|funding|funds received|receipts?)\b", _normal(label)) for label in labels):
                    warnings.append(f"Additional gain/income table at row {number} was not imported; review it separately as funding or income.")
                    section = "income"
                elif _header(values):
                    expected = {key: value for key, value in headers.items() if key in columns}
                    from openpyxl.utils import get_column_letter
                    if expected == {key: get_column_letter(index + 1) for key, index in columns.items()}:
                        if section:
                            warnings.append(f"Imported an additional matching ledger table at row {number}.")
                        section = None
                    else:
                        warnings.append(f"Additional table at row {number} has different columns and requires separate mapping; it was not imported.")
                        section = "unmapped"
                    skipped += 1
                    continue
            balance_label = not dated and any(re.search(r"\b(?:beginning|opening|closing|ending|added)\s+balance\b", label, re.I) for label in labels)
            if section or balance_label or any((BALANCE if dated else SUMMARY).search(label) for label in labels if label):
                skipped += 1
                continue
            # Formula caches are checked before interpreting absent values as blank rows.
            for field, index in columns.items():
                if raw_cells[index].data_type == "f" and values[index] is None:
                    raise DomainError(f"{sheet_name} row {number}: {field} formula has no cached value; recalculate and save in Excel", "formula_cache_missing", 422)
            if values[columns["date"]] in (None, ""):
                money_fields = [key for key in ("amount", "debit", "credit") if key in columns]
                require(not any(values[columns[key]] not in (None, "", 0) for key in money_fields),
                        f"{sheet_name} row {number}: transaction amount has no date", "invalid_date", 422)
                skipped += 1
                continue
            try:
                rows.append(_transaction(values, columns, actual_options, {"sheet": sheet_name, "row": number}, cached.epoch))
            except DomainError as exc:
                raise DomainError(f"{sheet_name} row {number}: {exc.message}", exc.code, exc.status) from exc
            require(len(rows) <= MAX_TRANSACTIONS, "Import is limited to 500 transactions", "file_limit", 413)
        require(rows, "No transaction rows were found with this mapping", "no_transactions", 422)
        from openpyxl.utils import get_column_letter
        if skipped:
            warnings.insert(0, f"Skipped {skipped} balance, total, summary or non-transaction rows.")
        return {"sheet": sheet_name, "headerRow": header_row, "mapping": {k: get_column_letter(v + 1) for k, v in columns.items()},
                "dateFormat": actual_options["dateFormat"], "defaultCurrency": actual_options.get("defaultCurrency", ""),
                "rows": rows, "warnings": warnings}
    finally:
        formulas.close()
        cached.close()


def parse_bank_csv(raw, options=None):
    options = options or {}
    _options(options)
    require(isinstance(raw, bytes) and 0 < len(raw) <= XLSX_MAX_BYTES, "Bank CSV must contain 1 byte to 2 MB", "file_limit", 413)
    try:
        text = raw.decode("utf-8-sig")
        delimiter = options.get("delimiter", ",")
        require(delimiter in {",", ";", "\t"}, "Choose comma, semicolon or tab delimiter")
        values = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    except (UnicodeDecodeError, csv.Error) as exc:
        raise DomainError("Bank CSV must be valid UTF-8 CSV", "invalid_csv", 422) from exc
    require(1 < len(values) <= MAX_ROWS and all(len(row) <= MAX_COLUMNS for row in values), "CSV exceeds row/column limits or has no records", "file_limit", 413)
    header_row = options.get("headerRow", 1)
    require(type(header_row) is int and 1 <= header_row < len(values), "Select an existing header row", "invalid_mapping", 422)
    headers = values[header_row - 1]
    columns = _mapping(options.get("mapping") or _header(headers), len(headers), bank=True)
    rows, skipped = [], 0
    for number, row in enumerate(values[header_row:], header_row + 1):
        if not any(v.strip() for v in row):
            continue
        require(len(row) == len(headers), f"CSV row {number} has the wrong number of columns", "invalid_csv", 422)
        if any(SUMMARY.search(value.strip()) for value in row):
            skipped += 1
            continue
        rows.append(_transaction(row, columns, options, {"row": number}, bank=True))
        require(len(rows) <= MAX_TRANSACTIONS, "Import is limited to 500 transactions", "file_limit", 413)
    return {"rows": rows, "warnings": [f"Skipped {skipped} summary rows."] if skipped else [], "recognized": True}


def _pdf_reader(raw):
    require(isinstance(raw, bytes) and 0 < len(raw) <= 5 * 1024 * 1024, "PDF must contain 1 byte to 5 MB", "file_limit", 413)
    require(raw.startswith(b"%PDF-"), "File is not a PDF", "invalid_document", 415)
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(raw))
        require(not reader.is_encrypted, "Encrypted PDFs are unsupported", "unsupported_document", 415)
        require(0 < len(reader.pages) <= 20, "PDF must have 1 to 20 pages", "file_limit", 413)
        return reader
    except DomainError:
        raise
    except Exception as exc:
        raise DomainError("PDF could not be read", "invalid_document", 415) from exc


def extract_pdf_pages(raw):
    reader = _pdf_reader(raw)
    pages = []
    try:
        for page in reader.pages:
            pages.append(page.extract_text() or "")
            require(sum(map(len, pages)) <= 100_000, "PDF text exceeds 100,000 characters", "file_limit", 413)
        require(all(p.strip() for p in pages), "Every PDF page needs extractable text; OCR is not supported", "ocr_unsupported", 415)
        return pages
    except DomainError:
        raise
    except Exception as exc:
        raise DomainError("PDF text could not be extracted", "invalid_document", 415) from exc


_BANK_MONEY = re.compile(r"(?<!\w)(?:\d{1,3}(?:,\d{3})+|\d+)\.\d{2}(?!\d)")


def parse_bank_pdf(raw):
    pages = extract_pdf_pages(raw)
    fallback = {"rows": [], "warnings": ["Bank layout is not recognised reliably; review the source and provide a mapped bank CSV or scoped extraction proposal."], "recognized": False, "textPages": pages}
    if not any("BANCA TRANSILVANIA" in page.upper() and "EXTRAS CONT" in page.upper() for page in pages):
        return fallback
    reader = _pdf_reader(raw)
    rows, opening, closing, credit_total, debit_total = [], None, None, Decimal(0), Decimal(0)
    currency = None
    for page_number, page in enumerate(reader.pages, 1):
        try:
            layout = page.extract_text(extraction_mode="layout")
        except Exception:
            return fallback
        lines = layout.splitlines()
        start = next((i for i, line in enumerate(lines) if re.search(r"\bData\b.*\bDescriere\b.*\bDebit\b.*\bCredit\b", line, re.I)), None)
        if start is None:
            return fallback
        header = lines[start]
        debit_column, credit_column = header.lower().find("debit"), header.lower().find("credit")
        currency_match = re.search(r"\b([A-Z]{3})\s+Cod IBAN\b", layout)
        if not currency_match:
            currency_match = re.search(r"Valuta\s*\n\s*([A-Z]{3})\b", layout)
        if not currency_match:
            # Currency is allowed on the line after the account line, before the table.
            currency_match = re.search(r"\b(RON|EUR|USD|CAD|GBP)\b", "\n".join(lines[max(0, start - 4):start]))
        if not currency_match or currency and currency != currency_match.group(1):
            return fallback
        currency = currency_match.group(1)
        active = None
        for line_number, line in enumerate(lines[start + 1:], start + 2):
            stripped = line.strip()
            numbers = list(_BANK_MONEY.finditer(line))
            if "SOLD ANTERIOR" in line.upper() and numbers:
                value = _number(numbers[-1].group(), "Opening balance", decimal_separator=".")
                if opening is None:
                    opening = value
                continue
            if "SOLD FINAL CONT" in line.upper() and numbers:
                closing = _number(numbers[-1].group(), "Closing balance", decimal_separator=".")
            summary = re.search(r"\b(?:RULAJ|SOLD FINAL|TOTAL DISPONIBIL|Fonduri proprii|Credit neutilizat)\b", line, re.I)
            dated = re.match(r"\s*(\d{2}/\d{2}/\d{4})\s+(.+)", line)
            if summary or "Acest extras" in line:
                if active is not None:
                    rows.append(active)
                active = None
                continue
            if dated:
                if active is not None:
                    rows.append(active)
                # A valid transaction amount must lie in a debit/credit table column.
                amounts = [m for m in numbers if m.end() >= debit_column - 3]
                if len(amounts) != 1:
                    return fallback
                found = amounts[0]
                if found.start() > 0 and line[found.start() - 1] in "-(":
                    return fallback
                direction = "credit" if found.end() > (debit_column + credit_column) / 2 else "debit"
                amount = _number(found.group(), "Transaction amount", decimal_separator=".")
                if amount <= 0:
                    return fallback
                description = line[dated.start(2):found.start()].strip()
                active = {"date": _day(dated.group(1)), "description": description, "reference": "", "amount": f"{amount:.2f}",
                          "direction": direction, "currency": currency, "category": "", "source": {"page": page_number, "row": line_number}}
                if direction == "credit":
                    credit_total += amount
                else:
                    debit_total += amount
            elif active is not None and stripped:
                if re.match(r"REF\s*:", stripped, re.I):
                    active["bankReference"] = stripped.split(":", 1)[1].strip()
                else:
                    active["description"] += " " + stripped
                adjustment = re.search(r"\b[A-Z]{2,8}-ADJ-[A-Z0-9-]+\b", stripped)
                if adjustment:
                    active["reference"] = adjustment.group()
                # The reference line ends this transaction's multi-line description.
                if re.match(r"REF\s*:", stripped, re.I):
                    if not active["reference"]:
                        active["reference"] = active["bankReference"]
                    rows.append(active)
                    active = None
            require(len(rows) <= MAX_TRANSACTIONS, "Bank import is limited to 500 transactions", "file_limit", 413)
        if active is not None:
            rows.append(active)
    if not rows:
        return fallback
    require(len(rows) <= MAX_TRANSACTIONS, "Bank import is limited to 500 transactions", "file_limit", 413)
    warnings = []
    if opening is not None and closing is not None:
        if opening + credit_total - debit_total != closing:
            return {**fallback, "warnings": ["Parsed transactions do not reconcile with the statement's opening and closing balances; review the source before importing."]}
    else:
        warnings.append("Opening/closing balance reconciliation was unavailable; review transaction completeness.")
    return {"rows": rows, "warnings": warnings, "recognized": True, "textPages": pages}


def _safe_cell(cell, value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # Preserve literal text, including = + - @, without formulas.


def _money_minor(value):
    require(type(value) is int, "Export amounts must be integer minor units", "invalid_financial_value", 422)
    return Decimal(value) / Decimal(100)


def export_financial_xlsx(report, entries, receipts):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    require(isinstance(report, dict) and isinstance(entries, list) and isinstance(receipts, list), "Invalid financial export")
    require(len(entries) <= MAX_TRANSACTIONS and len(receipts) <= MAX_TRANSACTIONS, "Export exceeds 500 rows per ledger", "file_limit", 413)
    book = Workbook()
    summary = book.active
    summary.title = "Financial Report"
    sheets = [(summary, ["Category", "Budget", "Actual", "Variance"]),
              (book.create_sheet("Ledger"), ["Date", "Description", "Reference", "Category", "Native amount", "Native currency", "Report amount", "Report currency", "Exchange rate (exact text)", "Source", "History", "Conversion basis"]),
              (book.create_sheet("Funding Receipts"), ["Date", "Description", "Reference", "Native amount", "Native currency", "Report amount", "Report currency", "Exchange rate (exact text)", "Source", "History", "Note", "Entered rate direction", "Entered rate (exact text)", "Receipt ID", "Automatic for drafts"])]
    marker = "Synthetic demonstration data" if report.get("synthetic") is True else "Financial report"
    for sheet, headers in sheets:
        _safe_cell(sheet.cell(1, 1), marker)
        for index, header in enumerate(headers, 1):
            _safe_cell(sheet.cell(3, index), header)
            sheet.cell(3, index).font = Font(bold=True, color="FFFFFF")
            sheet.cell(3, index).fill = PatternFill("solid", fgColor="0A7C78")
        sheet.freeze_panes = "A4"
    for number, row in enumerate(report.get("rows", []), 4):
        for column, value in enumerate([row["category"], _money_minor(row["budgetMinor"]), _money_minor(row["actualMinor"]), _money_minor(row["varianceMinor"])], 1):
            _safe_cell(summary.cell(number, column), value)
    report_currency = report.get("currency", report.get("reportCurrency", ""))
    total_row = len(report.get("rows", [])) + 4
    _safe_cell(summary.cell(total_row, 1), "Total")
    for column, field in enumerate(("budgetMinor", "actualMinor", "varianceMinor"), 2):
        _safe_cell(summary.cell(total_row, column), _money_minor(sum(row[field] for row in report.get("rows", []))))
        summary.cell(total_row, column).font = Font(bold=True)
    grant = report.get("grant") or {}
    metadata = [("Grant", report.get("grantName", grant.get("name", ""))), ("Reporting currency", report_currency),
                ("Period start", report.get("periodStart", "")), ("Period end", report.get("periodEnd", "")),
                ("Unresolved entries", report.get("unresolvedCount", 0)),
                ("Unresolved imports", report.get("unresolvedImportCount", 0)),
                ("Source warnings", '\n'.join(report.get("sourceWarnings", [])) or 'None')]
    for number, (label, value) in enumerate(metadata, 3):
        _safe_cell(summary.cell(number, 6), label)
        summary.cell(number, 6).font = Font(bold=True)
        _safe_cell(summary.cell(number, 7), value)
    summary.column_dimensions["F"].width = 25
    summary.column_dimensions["G"].width = 65
    for number, entry in enumerate(entries, 4):
        native = entry.get("sourceAmountMinor", entry.get("nativeAmountMinor", entry.get("amountMinor")))
        converted = entry.get("reportAmountMinor", entry.get("convertedAmountMinor"))
        rate = entry.get("fxRate", entry.get("rate", entry.get("exchangeRate")))
        basis = entry.get('conversion', {})
        basis_label = ('Automatic receipt: ' if basis.get('automaticReceipt') else 'Receipt: ') + basis['receiptId'] if basis.get('receiptId') else (
            'Weighted receipts: ' + ', '.join(basis['receiptIds']) if basis.get('receiptIds') else
            'Correction using original rate and cumulative rounding' if basis.get('roundingBasis') else entry.get('conversionMode', 'Recorded conversion'))
        values = [entry.get("date", ""), entry.get("description", ""), entry.get("reference", ""), entry.get("category", ""),
                  _money_minor(native), entry.get("nativeCurrency", entry.get("currency", "")), _money_minor(converted) if converted is not None else None,
                  entry.get("reportCurrency", report_currency), _rate_text(rate) if rate not in (None, "") else None, entry.get("source", {}), entry.get("history", []), basis_label]
        for column, value in enumerate(values, 1):
            _safe_cell(sheets[1][0].cell(number, column), value)
    for number, entry in enumerate(receipts, 4):
        converted, rate = entry.get("reportAmountMinor"), entry.get("rate")
        values = [entry.get("receivedDate", entry.get("date", "")), entry.get("description", ""), entry.get("reference", ""), _money_minor(entry.get("sourceAmountMinor", entry.get("amountMinor", entry.get("nativeAmountMinor")))),
                  entry.get("sourceCurrency", entry.get("currency", entry.get("nativeCurrency", ""))), _money_minor(converted) if converted is not None else None,
                  entry.get("reportCurrency", report_currency), _rate_text(rate) if rate not in (None, "") else None,
                  entry.get("source", {}), entry.get("history", []), entry.get("note", ""), entry.get("enteredDirection", ""),
                  str(entry["enteredRate"]) if entry.get("enteredRate") is not None else "", entry.get('id', ''), 'Yes' if entry.get('useAsDefault') else 'No']
        for column, value in enumerate(values, 1):
            _safe_cell(sheets[2][0].cell(number, column), value)
    from openpyxl.utils import get_column_letter
    for sheet, headers in sheets:
        filter_end = len(report.get("rows", [])) + 3 if sheet is summary else max(sheet.max_row, 3)
        sheet.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{filter_end}"
        for column in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 32 if "rate" in headers[column - 1].lower() else (24 if column != 2 else 45)
        sheet.row_dimensions[3].height = 32
        for cell in sheet[3]:
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        for row in sheet.iter_rows(min_row=4):
            sheet.row_dimensions[row[0].row].height = 34
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                if cell.data_type == "n":
                    cell.number_format = '#,##0.00;[Red](#,##0.00)'
    summary.column_dimensions['A'].width = 48
    for column in ('B', 'C', 'D'):
        summary.column_dimensions[column].width = 18
    sheets[1][0].column_dimensions['D'].width = 48
    sheets[1][0].column_dimensions['J'].width = 60
    sheets[1][0].column_dimensions['K'].width = 80
    sheets[1][0].column_dimensions['L'].width = 65
    sheets[2][0].column_dimensions['K'].width = 50
    sheets[2][0].column_dimensions['L'].width = 25
    sheets[2][0].column_dimensions['N'].width = 34
    summary['G7'].number_format = '0'
    summary['G8'].number_format = '0'
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def fill_report_template(raw, report, mapping=None):
    from openpyxl.cell.cell import MergedCell
    from openpyxl.utils import column_index_from_string
    book = _xlsx(raw, read_only=False)
    mapping = mapping or {}
    require(isinstance(mapping, dict) and isinstance(report, dict), "Template mapping and report must be objects")
    try:
        rate_direction = mapping.get("rateDirection", "report_per_source")
        require(isinstance(rate_direction, str) and rate_direction in {"report_per_source", "source_per_report"},
                "Choose report_per_source or source_per_report for the template rate direction", "invalid_mapping", 422)
        name = mapping.get("sheet", "Financial Report")
        require(isinstance(name, str) and name in book.sheetnames, "Select the report template sheet", "invalid_mapping", 422)
        sheet = book[name]
        category_cells = mapping.get("categories", mapping.get("rows", mapping.get("categoryRows")))
        report_rows = report.get("rows", [])
        require(isinstance(report_rows, list) and all(isinstance(r, dict) for r in report_rows), "Report category rows are required")
        actual_column = mapping.get("actualColumn", "G")
        require(isinstance(actual_column, str) and re.fullmatch(r"[A-Za-z]{1,2}", actual_column)
                and column_index_from_string(actual_column.upper()) <= MAX_COLUMNS, "Actual column must be a supported column letter", "invalid_mapping", 422)
        range_mapping = any(key in mapping for key in ("categoryColumn", "startRow", "endRow"))
        if range_mapping:
            require(category_cells is None, "Choose category cell mappings or a category row range", "invalid_mapping", 422)
            category_column, start, end = mapping.get("categoryColumn"), mapping.get("startRow"), mapping.get("endRow")
            require(isinstance(category_column, str) and re.fullmatch(r"[A-Za-z]{1,2}", category_column)
                    and column_index_from_string(category_column.upper()) <= MAX_COLUMNS
                    and category_column.upper() != actual_column.upper(), "Map separate category and actual columns", "invalid_mapping", 422)
            require(type(start) is int and type(end) is int and 1 <= start <= end <= MAX_ROWS,
                    "Map a valid category row range", "invalid_mapping", 422)
            category_cells = {}
            for number in range(start, end + 1):
                cell = sheet[f"{category_column.upper()}{number}"]
                require(cell.data_type != "f" and isinstance(cell.value, str) and cell.value.strip(),
                        "Every mapped category row needs a literal category label", "invalid_mapping", 422)
                label = cell.value.strip()
                require(label not in category_cells, "Template category labels must be unique", "invalid_mapping", 422)
                category_cells[label] = f"{actual_column.upper()}{number}"
        elif category_cells is None:
            require([_normal(sheet.cell(row, 2).value) for row in range(25, 33)] == [_normal(x) for x in GCC_CATEGORIES], "Unknown report layout; map each category's actual-value cell", "mapping_required", 422)
            category_cells = {str(sheet.cell(row, 2).value).strip(): f"G{row}" for row in range(25, 33)}
            mapping = {"totalCell": "G33", **mapping}
        require(isinstance(category_cells, dict) and bool(category_cells), "Map report categories to actual-value cells", "invalid_mapping", 422)
        require(all(isinstance(category, str) and category.strip() for category in category_cells), "Template categories must be nonempty labels", "invalid_mapping", 422)
        category_cells = {category: f"{actual_column}{target}" if type(target) is int else target for category, target in category_cells.items()}
        require(all(isinstance(row.get("category"), str) and row["category"].strip() for row in report_rows), "Report categories must be nonempty labels", "invalid_mapping", 422)
        categories = {_normal(row["category"]): row for row in report_rows}
        mapped_categories = {_normal(category): target for category, target in category_cells.items()}
        require(len(categories) == len(report_rows) and len(mapped_categories) == len(category_cells), "Report and template categories must be unique", "invalid_mapping", 422)
        require(set(categories).issubset(mapped_categories), "Template mapping must include every report category", "invalid_mapping", 422)
        writes = [(target, _money_minor(categories[category]["actualMinor"] if category in categories else 0)) for category, target in mapped_categories.items()]
        if mapping.get("totalCell"):
            writes.append((mapping["totalCell"], _money_minor(sum(row["actualMinor"] for row in report_rows))))
        if mapping.get("rateCell"):
            rate = report.get("rate", report.get("exchangeRate"))
            require(rate is not None, "An explicit rate is required for the mapped rate cell", "invalid_mapping", 422)
            display_rate = _rate_text(rate)
            if rate_direction == "source_per_report":
                with localcontext() as context:
                    context.prec = 40
                    inverse = (Decimal(1) / Decimal(display_rate)).quantize(Decimal("0.000000000000000001"), rounding=ROUND_HALF_UP)
                    display_rate = format(inverse.normalize(), "f")
            writes.append((mapping["rateCell"], display_rate))
        if mapping.get("periodStartCell"):
            writes.append((mapping["periodStartCell"], _day(report.get("periodStart"))))
        if mapping.get("periodEndCell"):
            writes.append((mapping["periodEndCell"], _day(report.get("periodEnd"))))
        coordinates = []
        for address, value in writes:
            require(isinstance(address, str) and re.fullmatch(r"[A-Za-z]{1,2}[1-9]\d{0,3}", address), "Map an individual cell address", "invalid_mapping", 422)
            cell = sheet[address.upper()]
            require(cell.row <= MAX_ROWS and cell.column <= MAX_COLUMNS and not isinstance(cell, MergedCell), "Mapped cell is outside supported bounds or inside a merged range", "invalid_mapping", 422)
            require(cell.coordinate not in coordinates, "Mapped output cells must not overlap", "invalid_mapping", 422)
            coordinates.append(cell.coordinate)
            _safe_cell(cell, value)  # Only explicitly mapped cells may replace formulas.
        output = io.BytesIO()
        book.save(output)
        return output.getvalue()
    finally:
        book.close()
